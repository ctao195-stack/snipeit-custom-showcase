#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Small internal web app for querying asset lifecycle history."""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urlencode, urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook


SOURCE_PATH = Path(r"D:\snipe-it\asset_activity.xlsx")
HOST = "0.0.0.0"
PORT = 18082
SHEET_NAME = "ActivityLog"
REPORT_DIR = Path(r"D:\资产周度汇总")
WEEKLY_REPORT_FILENAME = "资产周度汇总.xlsx"
SNIPEIT_ENV_PATH = Path(r"D:\snipe-it-excel-sync\.env.snipeit-excel")
SNIPEIT_COUNT_CACHE_SECONDS = 300
DATA_RELOAD_SECONDS = 5

STATUS_ORDER = [
    "入库",
    "领用",
    "借用中",
    "归还",
    "退租",
    "维修中",
    "维修完成已重新入库",
    "已换新",
    "删除资产",
]
STATUS_LABELS = set(STATUS_ORDER)

@dataclass
class ActivityRow:
    sync_time: str
    log_id: str
    action_time: datetime
    action: str
    status: str
    current_status: str
    category: str
    internal_id: str
    asset_tag: str
    asset_name: str
    serial: str
    operator: str
    target: str
    note: str
    meta: str


class LifecycleStore:
    def __init__(self, source_path: Path) -> None:
        self.source_path = source_path
        self.loaded_at = 0.0
        self.mtime = 0.0
        self.rows: list[ActivityRow] = []
        self.grouped: dict[str, list[ActivityRow]] = {}
        self.index: dict[str, str] = {}
        self.current_asset_count_cache: dict[str, Any] = {
            "count": None,
            "loadedAt": 0.0,
            "error": "",
        }

    def refresh_if_needed(self) -> None:
        if not self.source_path.exists():
            self.rows = []
            self.grouped = {}
            self.index = {}
            self.loaded_at = time.time()
            self.mtime = 0
            return
        mtime = self.source_path.stat().st_mtime
        if self.rows and mtime == self.mtime and time.time() - self.loaded_at < DATA_RELOAD_SECONDS:
            return
        self.rows = load_rows(self.source_path)
        self.grouped = group_by_asset(self.rows)
        self.index = {}
        for asset_tag, rows in self.grouped.items():
            latest = rows[-1]
            values = {
                normalize_asset_key(asset_tag),
                normalize_asset_key(latest.serial),
                normalize_asset_key(latest.asset_name),
                normalize_asset_key(latest.internal_id),
            }
            for value in values:
                if value:
                    self.index.setdefault(value, asset_tag)
        self.loaded_at = time.time()
        self.mtime = mtime

    def search(self, query: str) -> dict[str, Any]:
        self.refresh_if_needed()
        clean_query = normalize_asset_key(query)
        if not clean_query:
            return {
                "ok": False,
                "message": "请输入资产编号、序列号或资产名称关键字。",
                "query": query,
                "suggestions": self.suggestions(""),
            }

        exact_asset = self.index.get(clean_query)
        if exact_asset:
            return {"ok": True, "asset": self.asset_payload(exact_asset)}

        matches = self.find_matches(clean_query)
        if len(matches) == 1:
            return {"ok": True, "asset": self.asset_payload(matches[0])}
        return {
            "ok": False,
            "message": "没有找到完全匹配的资产，下面是可能相关的结果。",
            "query": query,
            "suggestions": [self.summary_payload(asset) for asset in matches[:30]],
        }

    def find_matches(self, clean_query: str) -> list[str]:
        scored: list[tuple[int, str]] = []
        for asset_tag, rows in self.grouped.items():
            latest = rows[-1]
            fields = [
                normalize_asset_key(asset_tag),
                normalize_asset_key(latest.serial),
                normalize_asset_key(latest.asset_name),
                normalize_asset_key(latest.category),
                normalize_asset_key(latest.target),
            ]
            score = 0
            for field in fields:
                if not field:
                    continue
                if field == clean_query:
                    score = max(score, 100)
                elif field.startswith(clean_query):
                    score = max(score, 80)
                elif clean_query in field:
                    score = max(score, 50)
            if score:
                scored.append((score, asset_tag))
        scored.sort(key=lambda item: (-item[0], item[1]))
        return [asset_tag for _, asset_tag in scored]

    def suggestions(self, query: str) -> list[dict[str, Any]]:
        self.refresh_if_needed()
        clean_query = normalize_asset_key(query)
        if clean_query:
            assets = self.find_matches(clean_query)[:20]
        else:
            assets = sorted(
                self.grouped.keys(),
                key=lambda asset: self.grouped[asset][-1].action_time,
                reverse=True,
            )[:20]
        return [self.summary_payload(asset) for asset in assets]

    def dashboard(self) -> dict[str, Any]:
        self.refresh_if_needed()
        latest_rows = [rows[-1] for rows in self.grouped.values() if rows]
        status_counts = Counter(row.status for row in latest_rows)
        category_counts = Counter(row.category or "未填写" for row in latest_rows)

        now = datetime.now()
        week_start = start_of_week(now)
        weekly_rows = [row for row in self.rows if week_start <= row.action_time <= now]
        recent_rows = sorted(self.rows, key=lambda row: row.action_time, reverse=True)[:30]

        current_asset_info = self.current_asset_count()
        return {
            "status": self.status(),
            "overview": {
                "currentAssetCount": current_asset_info.get("count"),
                "currentAssetCountError": current_asset_info.get("error", ""),
                "historicalAssetCount": len(self.grouped),
                "assetCount": len(self.grouped),
                "rowCount": len(self.rows),
                "weeklyCount": len(weekly_rows),
                "weekStart": week_start.strftime("%Y-%m-%d"),
                "weekEnd": (week_start + timedelta(days=6)).strftime("%Y-%m-%d"),
            },
            "statusCounts": [
                {"status": status, "count": status_counts.get(status, 0)}
                for status in STATUS_ORDER
                if status_counts.get(status, 0)
            ]
            + [
                {"status": status, "count": count}
                for status, count in sorted(status_counts.items())
                if status not in STATUS_LABELS
            ],
            "categoryCounts": [
                {"category": category, "count": count}
                for category, count in category_counts.most_common(8)
            ],
            "weeklyStatusCounts": [
                {"status": status, "count": count}
                for status, count in Counter(row.status for row in weekly_rows).most_common()
            ],
            "weeklyActionCounts": [
                {"action": action or "未填写", "count": count}
                for action, count in Counter(row.action for row in weekly_rows).most_common()
            ],
            "recent": [row_payload(row) for row in recent_rows],
            "weeklyReport": self.weekly_report(),
        }

    def current_asset_count(self) -> dict[str, Any]:
        now = time.time()
        cached_at = float(self.current_asset_count_cache.get("loadedAt") or 0)
        if cached_at and now - cached_at < SNIPEIT_COUNT_CACHE_SECONDS:
            return self.current_asset_count_cache

        try:
            count = fetch_snipeit_current_asset_count()
            self.current_asset_count_cache = {
                "count": count,
                "loadedAt": now,
                "error": "",
            }
        except Exception as exc:
            self.current_asset_count_cache = {
                "count": None,
                "loadedAt": now,
                "error": str(exc),
            }
        return self.current_asset_count_cache

    def weekly_report(self) -> dict[str, Any]:
        self.refresh_if_needed()
        now = datetime.now()
        week_start = start_of_week(now)
        previous_start = week_start - timedelta(days=7)
        previous_end = now - timedelta(days=7)
        rows = [
            row for row in self.rows if week_start <= row.action_time <= now
        ]
        previous_rows = [
            row for row in self.rows if previous_start <= row.action_time <= previous_end
        ]
        asset_count = len({row.asset_tag for row in rows})
        status_items = ordered_count_items(Counter(row.status for row in rows), "status")
        action_items = count_items(Counter(row.action or "未填写" for row in rows), "action")
        category_items = count_items(Counter(row.category or "未填写" for row in rows), "category")
        week_label = iso_week_label(now)
        delta = len(rows) - len(previous_rows)
        delta_text = change_text(delta)
        status_text = count_summary(status_items, "status")
        action_text = count_summary(action_items, "action")
        category_text = count_summary(category_items[:5], "category")
        report_lines = [
            f"资产周报（{week_label}，{week_start:%Y-%m-%d} 至 {now:%Y-%m-%d}）",
            f"1. 本周资产系统共记录操作 {len(rows)} 条，涉及资产 {asset_count} 台，{delta_text}。",
            f"2. 按状态结果统计：{status_text}。",
            f"3. 按操作类型统计：{action_text}。",
            f"4. 涉及资产分类：{category_text}。",
        ]
        latest_file = latest_weekly_report_file()
        return {
            "weekLabel": week_label,
            "weekStart": week_start.strftime("%Y-%m-%d"),
            "weekEnd": now.strftime("%Y-%m-%d"),
            "operationCount": len(rows),
            "assetCount": asset_count,
            "previousOperationCount": len(previous_rows),
            "delta": delta,
            "statusCounts": status_items,
            "actionCounts": action_items,
            "categoryCounts": category_items,
            "copyText": "\n".join(report_lines),
            "recent": [row_payload(row) for row in sorted(rows, key=lambda item: item.action_time, reverse=True)[:80]],
            "download": {
                "exists": latest_file is not None,
                "url": "/download/weekly-report" if latest_file else "",
                "fileName": latest_file.name if latest_file else "",
                "modifiedAt": datetime.fromtimestamp(latest_file.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
                if latest_file
                else "",
            },
        }

    def person_search(self, query: str) -> dict[str, Any]:
        self.refresh_if_needed()
        clean_query = normalize_asset_key(query)
        if not clean_query:
            return {
                "ok": False,
                "message": "请输入员工姓名、部门关键字或使用对象关键字。",
                "query": query,
                "holdings": [],
                "history": [],
            }

        holdings = []
        history = []
        for asset_tag, rows in self.grouped.items():
            if not rows:
                continue
            latest = rows[-1]
            if latest.status in {"领用", "借用中"} and text_contains(latest.target, clean_query):
                holdings.append(self.summary_payload(asset_tag))
            for row in rows:
                if (
                    text_contains(row.target, clean_query)
                    or text_contains(row.note, clean_query)
                    or text_contains(row.operator, clean_query)
                ):
                    history.append(row)

        holdings.sort(key=lambda item: item.get("latestTime", ""), reverse=True)
        history.sort(key=lambda row: row.action_time, reverse=True)
        return {
            "ok": bool(holdings or history),
            "message": "" if holdings or history else "没有找到与该人员/对象相关的资产记录。",
            "query": query,
            "currentCount": len(holdings),
            "activityCount": len(history),
            "holdings": holdings[:200],
            "history": [row_payload(row) for row in history[:300]],
        }

    def asset_payload(self, asset_tag: str) -> dict[str, Any]:
        rows = self.grouped.get(asset_tag, [])
        if not rows:
            return {}
        latest = rows[-1]
        first = rows[0]
        counts = Counter(row.status for row in rows)
        issues = asset_issues(rows)
        return {
            "assetTag": latest.asset_tag,
            "assetName": latest.asset_name or first.asset_name,
            "category": latest.category or first.category,
            "serial": latest.serial or first.serial,
            "currentStatus": latest.status,
            "currentHolder": current_holder(latest),
            "firstTime": format_dt(first.action_time),
            "latestTime": format_dt(latest.action_time),
            "totalCount": len(rows),
            "inCount": counts.get("入库", 0),
            "checkoutCount": counts.get("领用", 0) + counts.get("借用中", 0),
            "checkinCount": counts.get("归还", 0),
            "repairCount": counts.get("维修中", 0) + counts.get("维修完成已重新入库", 0),
            "disposeCount": counts.get("退租", 0) + counts.get("删除资产", 0),
            "latestNote": latest.note,
            "issues": issues,
            "history": [row_payload(row) for row in rows],
        }

    def summary_payload(self, asset_tag: str) -> dict[str, Any]:
        rows = self.grouped.get(asset_tag, [])
        if not rows:
            return {}
        latest = rows[-1]
        return {
            "assetTag": latest.asset_tag,
            "assetName": latest.asset_name,
            "category": latest.category,
            "serial": latest.serial,
            "currentStatus": latest.status,
            "currentHolder": current_holder(latest),
            "latestTime": format_dt(latest.action_time),
        }

    def status(self) -> dict[str, Any]:
        self.refresh_if_needed()
        return {
            "source": str(self.source_path),
            "exists": self.source_path.exists(),
            "assetCount": len(self.grouped),
            "rowCount": len(self.rows),
            "loadedAt": datetime.fromtimestamp(self.loaded_at).strftime("%Y-%m-%d %H:%M:%S")
            if self.loaded_at
            else "",
            "sourceMtime": datetime.fromtimestamp(self.mtime).strftime("%Y-%m-%d %H:%M:%S")
            if self.mtime
            else "",
        }


def load_rows(path: Path) -> list[ActivityRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
    header_map = {
        safe_text(sheet.cell(1, col).value): col - 1
        for col in range(1, sheet.max_column + 1)
        if safe_text(sheet.cell(1, col).value)
    }

    rows: list[ActivityRow] = []
    for raw in sheet.iter_rows(min_row=2, max_col=sheet.max_column, values_only=True):
        row = list(raw)
        action_time = parse_datetime(get_value(row, header_map, "操作时间"))
        asset_tag = safe_text(get_value(row, header_map, "资产编号"))
        if not action_time or not asset_tag:
            continue
        status = clean_status(
            get_value(row, header_map, "状态标签", "操作后状态", "资产状态", "当前资产状态")
        )
        current_status = clean_status(
            get_value(row, header_map, "_同步时当前状态", "同步时当前状态", "当前资产状态")
        )
        rows.append(
            ActivityRow(
                sync_time=safe_text(get_value(row, header_map, "同步时间")),
                log_id=safe_text(get_value(row, header_map, "_日志ID", "日志ID")),
                action_time=action_time,
                action=safe_text(get_value(row, header_map, "操作动作", "操作类型", "操作类型(中文)")),
                status=status or current_status or "未分类",
                current_status=current_status,
                category=safe_text(get_value(row, header_map, "资产分类", "资产类型")) or "未填写",
                internal_id=safe_text(
                    get_value(row, header_map, "_Snipe-IT内部ID", "Snipe-IT内部ID", "资产ID")
                ),
                asset_tag=asset_tag,
                asset_name=safe_text(get_value(row, header_map, "资产名称")),
                serial=safe_text(get_value(row, header_map, "序列号")),
                operator=safe_text(get_value(row, header_map, "操作人")),
                target=safe_text(get_value(row, header_map, "使用人/对象", "对象/领用人")),
                note=safe_text(get_value(row, header_map, "备注")),
                meta=safe_text(get_value(row, header_map, "_变更详情(JSON)", "变更详情(JSON)")),
            )
        )
    workbook.close()
    rows.sort(key=lambda item: (item.asset_tag, item.action_time, item.log_id))
    return rows


def get_value(row: list[Any], header_map: dict[str, int], *names: str) -> Any:
    for name in names:
        index = header_map.get(name)
        if index is not None and index < len(row):
            return row[index]
    return ""


def group_by_asset(rows: list[ActivityRow]) -> dict[str, list[ActivityRow]]:
    grouped: dict[str, list[ActivityRow]] = defaultdict(list)
    for row in rows:
        grouped[row.asset_tag].append(row)
    return {key: sorted(value, key=lambda item: (item.action_time, item.log_id)) for key, value in grouped.items()}


def row_payload(row: ActivityRow) -> dict[str, Any]:
    return {
        "assetTag": row.asset_tag,
        "actionTime": format_dt(row.action_time),
        "status": row.status,
        "action": row.action,
        "category": row.category,
        "assetName": row.asset_name,
        "serial": row.serial,
        "target": row.target,
        "operator": row.operator,
        "note": row.note,
        "syncTime": row.sync_time,
    }


def asset_issues(rows: list[ActivityRow]) -> list[str]:
    latest = rows[-1]
    statuses = [row.status for row in rows]
    issues = []
    if latest.status == "借用中":
        issues.append("当前处于借用中")
    if latest.status == "维修中":
        issues.append("维修中未完成")
    if "未分类" in statuses:
        issues.append("存在未分类记录")
    if latest.status == "删除资产":
        issues.append("资产已删除")
    if latest.status == "退租":
        issues.append("资产已退租")
    if has_activity_after_disposal(rows):
        issues.append("退租/删除后又出现操作")
    return dedupe(issues)


def has_activity_after_disposal(rows: list[ActivityRow]) -> bool:
    seen_disposal = False
    for row in rows:
        if seen_disposal and row.status not in {"退租", "删除资产"}:
            return True
        if row.status in {"退租", "删除资产"}:
            seen_disposal = True
    return False


def current_holder(row: ActivityRow) -> str:
    if row.status in {"领用", "借用中"}:
        return row.target
    return ""


def parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = safe_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def clean_status(value: Any) -> str:
    text = safe_text(value).strip().rstrip(".。…").strip()
    return text if text in STATUS_LABELS else (text or "")


def normalize_asset_key(value: Any) -> str:
    text = safe_text(value).upper()
    text = re.sub(r"\s+", "", text)
    text = text.replace("\u3000", "").replace("\ufeff", "")
    return text


def text_contains(value: Any, clean_query: str) -> bool:
    return clean_query in normalize_asset_key(value)


def load_env_file(path: Path) -> dict[str, str]:
    env: dict[str, str] = {}
    if not path.exists():
        return env
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip()
    return env


def fetch_snipeit_current_asset_count() -> int:
    env = load_env_file(SNIPEIT_ENV_PATH)
    base_url = env.get("SNIPEIT_URL", "").rstrip("/")
    token = env.get("SNIPEIT_API_TOKEN", "")
    if not base_url or not token:
        raise RuntimeError("Snipe-IT API 配置缺失")
    url = f"{base_url}/api/v1/hardware?{urlencode({'limit': 1, 'offset': 0})}"
    request = Request(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/json",
            "User-Agent": "asset-lifecycle-web/1.0",
        },
    )
    with urlopen(request, timeout=15) as response:
        data = json.loads(response.read().decode("utf-8"))
    total = data.get("total")
    if not isinstance(total, int):
        raise RuntimeError("Snipe-IT API 未返回 total")
    return total


def start_of_week(value: datetime) -> datetime:
    return value.replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(
        days=value.weekday()
    )


def iso_week_label(value: datetime) -> str:
    year, week, _ = value.isocalendar()
    return f"{year}-W{week:02d}"


def ordered_count_items(counter: Counter[str], key_name: str) -> list[dict[str, Any]]:
    items = [
        {key_name: status, "count": counter.get(status, 0)}
        for status in STATUS_ORDER
        if counter.get(status, 0)
    ]
    items.extend(
        {key_name: status, "count": count}
        for status, count in counter.most_common()
        if status not in STATUS_LABELS
    )
    return items


def count_items(counter: Counter[str], key_name: str) -> list[dict[str, Any]]:
    return [{key_name: name, "count": count} for name, count in counter.most_common()]


def count_summary(items: list[dict[str, Any]], key_name: str) -> str:
    if not items:
        return "暂无记录"
    return "，".join(f"{item.get(key_name) or '未填写'} {item.get('count', 0)} 条" for item in items)


def change_text(delta: int) -> str:
    if delta > 0:
        return f"较上周同期增加 {delta} 条"
    if delta < 0:
        return f"较上周同期减少 {abs(delta)} 条"
    return "较上周同期持平"


def latest_weekly_report_file() -> Path | None:
    fixed_path = REPORT_DIR / WEEKLY_REPORT_FILENAME
    if fixed_path.exists():
        return fixed_path
    if not REPORT_DIR.exists():
        return None
    files = [path for path in REPORT_DIR.glob("*.xlsx") if path.is_file()]
    if not files:
        return None
    return max(files, key=lambda path: path.stat().st_mtime)


def generate_weekly_report_file() -> Path | None:
    base_dir = Path(__file__).parent
    candidates = [
        base_dir / "generate_weekly_asset_report.py",
        base_dir.parent / "周度资产汇总" / "generate_weekly_asset_report.py",
    ]
    script_path = next((path for path in candidates if path.exists()), candidates[0])
    if not script_path.exists():
        return latest_weekly_report_file()
    try:
        subprocess.run(
            [sys.executable, str(script_path)],
            cwd=str(script_path.parent),
            check=False,
            timeout=120,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    except Exception:
        return latest_weekly_report_file()
    return latest_weekly_report_file()


def format_dt(value: datetime) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S")


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def dedupe(values: list[str]) -> list[str]:
    seen = set()
    result = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


HTML = r"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>资产管理看板</title>
  <style>
    :root {
      --primary: #17365d;
      --primary-2: #23527c;
      --bg: #f4f7fb;
      --panel: #ffffff;
      --field: #ffffff;
      --table-bg: #ffffff;
      --table-alt: #f8fbfd;
      --line: #d8e1ea;
      --text: #17202a;
      --muted: #657386;
      --soft: #eef5fa;
      --shadow: rgba(23, 54, 93, .10);
      --green: #d9ead3;
      --blue: #ddebf7;
      --purple: #e4dfec;
      --yellow: #fff2cc;
      --orange: #fce4d6;
      --gray: #e7e6e6;
      --red: #f4cccc;
    }
    body[data-theme="dark"] {
      --primary: #10243d;
      --primary-2: #183a5f;
      --bg: #101820;
      --panel: #172331;
      --field: #111c27;
      --table-bg: #14202d;
      --table-alt: #182738;
      --line: #314458;
      --text: #e7edf5;
      --muted: #a9b7c8;
      --soft: #213448;
      --shadow: rgba(0, 0, 0, .28);
      --green: #254331;
      --blue: #1f3e5b;
      --purple: #3a3151;
      --yellow: #4b3f1f;
      --orange: #4b3024;
      --gray: #33404c;
      --red: #5a2b2f;
      color-scheme: dark;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      font-family: "Microsoft YaHei", "PingFang SC", Arial, sans-serif;
      color: var(--text);
      background: var(--bg);
    }
    header {
      background: var(--primary);
      color: #fff;
      padding: 18px 28px;
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 16px;
    }
    header h1 {
      margin: 0;
      font-size: 22px;
      font-weight: 700;
    }
    .header-side {
      display: flex;
      align-items: center;
      gap: 12px;
    }
    header .meta {
      font-size: 13px;
      opacity: .92;
      text-align: right;
      line-height: 1.6;
    }
    .theme-toggle {
      height: 36px;
      border: 1px solid rgba(255, 255, 255, .38);
      background: rgba(255, 255, 255, .12);
      color: #fff;
      padding: 0 12px;
      min-width: 92px;
    }
    .theme-toggle:hover { background: rgba(255, 255, 255, .2); }
    main {
      max-width: 1560px;
      margin: 0 auto;
      padding: 22px 24px 38px;
    }
    .grid {
      display: grid;
      gap: 14px;
    }
    .overview {
      grid-template-columns: repeat(5, minmax(0, 1fr));
      margin-bottom: 14px;
    }
    .workbench {
      grid-template-columns: 420px minmax(0, 1fr);
      align-items: start;
      margin-bottom: 14px;
    }
    .query-stack {
      display: grid;
      gap: 14px;
    }
    .two-col {
      grid-template-columns: minmax(0, 1.1fr) minmax(0, .9fr);
      margin-top: 14px;
    }
    .lower-grid {
      grid-template-columns: minmax(0, 1.15fr) minmax(0, .85fr);
      margin-top: 14px;
      align-items: start;
    }
    .no-top-margin { margin-top: 0; }
    .top-space { margin-top: 14px; }
    .stats-grid {
      grid-template-columns: repeat(3, minmax(0, 1fr));
      align-items: start;
    }
    .panel, .kpi {
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      overflow: hidden;
      box-shadow: 0 1px 2px var(--shadow);
    }
    .panel h2 {
      margin: 0;
      padding: 12px 16px;
      font-size: 15px;
      color: #fff;
      background: var(--primary-2);
    }
    .panel-body {
      padding: 14px 16px 16px;
    }
    .kpi {
      padding: 13px 16px;
      min-height: 78px;
    }
    .kpi .label, .mini-label {
      font-size: 12px;
      color: var(--muted);
      margin-bottom: 7px;
    }
    .kpi .value {
      font-size: 22px;
      font-weight: 800;
      line-height: 1.25;
      word-break: break-word;
    }
    .input-row {
      display: grid;
      grid-template-columns: minmax(190px, 1fr) auto auto;
      gap: 10px;
    }
    input {
      height: 42px;
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 0 13px;
      font-size: 15px;
      outline: none;
      color: var(--text);
      background: var(--field);
    }
    input::placeholder { color: var(--muted); }
    input:focus {
      border-color: var(--primary-2);
      box-shadow: 0 0 0 3px rgba(35, 82, 124, .12);
    }
    button {
      height: 42px;
      border: 0;
      border-radius: 6px;
      padding: 0 15px;
      font-size: 14px;
      font-weight: 700;
      cursor: pointer;
      background: var(--primary-2);
      color: #fff;
      white-space: nowrap;
    }
    button.secondary {
      background: var(--soft);
      color: var(--primary);
      border: 1px solid var(--line);
    }
    body[data-theme="dark"] button.secondary {
      color: var(--text);
    }
    button.small {
      height: 32px;
      padding: 0 10px;
      font-size: 12px;
    }
    button:disabled {
      opacity: .55;
      cursor: not-allowed;
    }
    .hint {
      color: var(--muted);
      font-size: 13px;
      line-height: 1.6;
      margin-top: 10px;
    }
    .toolbar {
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      align-items: center;
      margin: 12px 0 0;
    }
    .report-text {
      width: 100%;
      min-height: 172px;
      resize: vertical;
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 12px 13px;
      font-family: "Microsoft YaHei", "PingFang SC", Arial, sans-serif;
      font-size: 14px;
      line-height: 1.7;
      color: var(--text);
      background: var(--field);
      outline: none;
    }
    .chip-grid {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(150px, 1fr));
      gap: 9px;
    }
    .chip {
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 10px 12px;
      background: var(--table-bg);
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 8px;
      min-height: 42px;
    }
    .chip strong { font-size: 17px; }
    .badge {
      display: inline-block;
      padding: 4px 9px;
      border-radius: 999px;
      border: 1px solid rgba(0,0,0,.06);
      background: var(--soft);
      font-weight: 700;
      line-height: 1.35;
      white-space: nowrap;
    }
    .tone-green { background: var(--green); }
    .tone-blue { background: var(--blue); }
    .tone-purple { background: var(--purple); }
    .tone-yellow { background: var(--yellow); }
    .tone-orange { background: var(--orange); }
    .tone-gray { background: var(--gray); }
    .tone-red { background: var(--red); }
    .tone-soft { background: var(--soft); }
    .suggestions {
      margin-top: 10px;
      border: 1px solid var(--line);
      border-radius: 7px;
      overflow: hidden;
      background: var(--table-bg);
      max-height: 260px;
      overflow-y: auto;
    }
    .asset-card {
      padding: 12px 14px;
      border-bottom: 1px solid var(--line);
      cursor: pointer;
    }
    .asset-card:last-child { border-bottom: 0; }
    .asset-card:hover { background: var(--soft); }
    .asset-card strong {
      display: block;
      font-size: 14px;
      margin-bottom: 4px;
    }
    .asset-card span {
      display: block;
      color: var(--muted);
      font-size: 12px;
      line-height: 1.5;
    }
    .summary {
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 10px;
      margin-bottom: 14px;
    }
    .mini-card {
      border: 1px solid var(--line);
      border-radius: 7px;
      background: var(--table-bg);
      padding: 11px 12px;
      min-height: 70px;
    }
    .mini-card .value {
      font-size: 16px;
      font-weight: 800;
      word-break: break-word;
    }
    .issues {
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
      margin-bottom: 14px;
    }
    .issue {
      background: var(--yellow);
      border: 1px solid #e7d78f;
      border-radius: 999px;
      padding: 5px 10px;
      font-size: 13px;
      font-weight: 700;
    }
    .issue-ok {
      background: var(--green);
      border-color: #87a977;
    }
    .table-wrap {
      max-height: 420px;
      overflow: auto;
      border: 1px solid var(--line);
      border-radius: 8px;
      background: var(--table-bg);
    }
    .result-table { max-height: 520px; }
    .short-table { max-height: 300px; }
    table {
      width: 100%;
      border-collapse: collapse;
      background: var(--table-bg);
      font-size: 13px;
    }
    th {
      position: sticky;
      top: 0;
      z-index: 1;
      background: var(--primary-2);
      color: #fff;
      text-align: left;
      padding: 9px 8px;
      border: 1px solid var(--line);
      white-space: nowrap;
    }
    td {
      padding: 9px 8px;
      border: 1px solid var(--line);
      vertical-align: top;
      line-height: 1.5;
    }
    tr:nth-child(even) td { background: var(--table-alt); }
    .muted { color: var(--muted); }
    .section-space { margin-top: 14px; }
    .subhead {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 10px;
      margin: 12px 0 8px;
    }
    .subhead .mini-label { margin-bottom: 0; }
    .hidden { display: none; }
    @media (max-width: 1100px) {
      .overview, .summary { grid-template-columns: repeat(2, minmax(0, 1fr)); }
      .workbench, .two-col, .lower-grid, .stats-grid { grid-template-columns: 1fr; }
    }
    @media (max-width: 720px) {
      header { display: block; }
      header .meta { text-align: left; margin-top: 8px; }
      main { padding: 16px 12px 28px; }
      .overview, .summary { grid-template-columns: 1fr; }
      .input-row { grid-template-columns: 1fr; }
    }
  </style>
</head>
<body>
  <header>
    <h1>资产管理看板</h1>
    <div class="header-side">
      <button class="theme-toggle" id="themeToggleBtn" type="button">自动外观</button>
      <div class="meta" id="meta">正在读取数据...</div>
    </div>
  </header>
  <main>
    <section class="grid overview" id="overview"></section>

    <section class="grid workbench">
      <div class="query-stack">
        <div class="panel">
          <h2>资产查询</h2>
          <div class="panel-body">
            <div class="input-row">
              <input id="assetQuery" placeholder="资产编号 / 序列号 / 名称关键字" autocomplete="off" />
              <button id="assetSearchBtn">查询</button>
              <button class="secondary" id="assetClearBtn">清空</button>
            </div>
            <div class="hint" id="assetNotice">复制编号里带换行或空格也可以自动处理。</div>
            <div id="suggestions" class="suggestions hidden"></div>
          </div>
        </div>
        <div class="panel">
          <h2>人员查询</h2>
          <div class="panel-body">
            <div class="input-row">
              <input id="personQuery" placeholder="员工姓名 / 部门 / 使用对象关键字" autocomplete="off" />
              <button id="personSearchBtn">查询</button>
              <button class="secondary" id="personClearBtn">清空</button>
            </div>
            <div class="hint" id="personNotice">可查看当前领用资产和相关历史记录。</div>
          </div>
        </div>
      </div>

      <div class="panel" id="weeklyReportPanel">
        <h2>周报区</h2>
        <div class="panel-body">
          <div id="weeklyReportKpis" class="summary"></div>
          <textarea id="weeklyReportText" class="report-text" readonly>正在生成周报...</textarea>
          <div class="toolbar">
            <button id="copyWeeklyReportBtn">复制周报文字</button>
            <button class="secondary" id="downloadWeeklyReportBtn">下载周报 Excel</button>
            <span class="hint" id="weeklyReportNotice"></span>
          </div>
        </div>
      </div>
    </section>

    <section class="panel hidden" id="assetResult">
      <h2>资产生命周期</h2>
      <div class="panel-body">
        <div id="assetSummary" class="summary"></div>
        <div id="assetIssues" class="issues"></div>
        <div class="table-wrap result-table">
          <table>
            <thead>
              <tr>
                <th>操作时间</th>
                <th>状态</th>
                <th>动作</th>
                <th>使用人/对象</th>
                <th>操作人</th>
                <th>备注</th>
              </tr>
            </thead>
            <tbody id="assetHistory"></tbody>
          </table>
        </div>
      </div>
    </section>

    <section class="panel hidden" id="personResult">
      <h2>人员资产</h2>
      <div class="panel-body">
        <div id="personSummary" class="summary"></div>
        <div class="grid two-col no-top-margin">
          <div>
            <div class="mini-label">当前领用资产</div>
            <div class="table-wrap short-table">
              <table>
                <thead>
                  <tr>
                    <th>资产编号</th>
                    <th>状态</th>
                    <th>资产名称</th>
                    <th>分类</th>
                    <th>最近操作</th>
                  </tr>
                </thead>
                <tbody id="personHoldings"></tbody>
              </table>
            </div>
          </div>
          <div>
            <div class="mini-label">相关历史</div>
            <div class="table-wrap short-table">
              <table>
                <thead>
                  <tr>
                    <th>资产编号</th>
                    <th>操作时间</th>
                    <th>状态</th>
                    <th>对象</th>
                  </tr>
                </thead>
                <tbody id="personHistory"></tbody>
              </table>
            </div>
          </div>
        </div>
      </div>
    </section>

    <section class="grid lower-grid">
      <div class="panel">
        <h2>最近操作</h2>
        <div class="panel-body">
          <div class="table-wrap">
            <table>
              <thead>
                <tr>
                  <th>资产编号</th>
                  <th>时间</th>
                  <th>状态</th>
                  <th>对象</th>
                </tr>
              </thead>
              <tbody id="recent"></tbody>
            </table>
          </div>
        </div>
      </div>
    </section>

    <section class="panel top-space">
      <h2>本周统计</h2>
      <div class="panel-body">
        <div class="grid stats-grid no-top-margin">
          <div>
            <div class="mini-label">本周状态结果</div>
            <div id="weeklyStatus" class="chip-grid"></div>
          </div>
          <div>
            <div class="mini-label">本周操作类型</div>
            <div id="weeklyAction" class="chip-grid"></div>
          </div>
          <div>
            <div class="mini-label">涉及资产分类</div>
            <div id="weeklyCategory" class="chip-grid"></div>
          </div>
        </div>
      </div>
    </section>
  </main>
  <script>
    const metaEl = document.getElementById('meta');
    const themeToggleBtn = document.getElementById('themeToggleBtn');
    const overviewEl = document.getElementById('overview');
    const assetQuery = document.getElementById('assetQuery');
    const assetNotice = document.getElementById('assetNotice');
    const suggestionsEl = document.getElementById('suggestions');
    const assetResult = document.getElementById('assetResult');
    const assetSummary = document.getElementById('assetSummary');
    const assetIssues = document.getElementById('assetIssues');
    const assetHistory = document.getElementById('assetHistory');
    const personQuery = document.getElementById('personQuery');
    const personNotice = document.getElementById('personNotice');
    const personResult = document.getElementById('personResult');
    const personSummary = document.getElementById('personSummary');
    const personHoldings = document.getElementById('personHoldings');
    const personHistory = document.getElementById('personHistory');
    const recentEl = document.getElementById('recent');
    const weeklyStatusEl = document.getElementById('weeklyStatus');
    const weeklyActionEl = document.getElementById('weeklyAction');
    const weeklyCategoryEl = document.getElementById('weeklyCategory');
    const weeklyReportKpis = document.getElementById('weeklyReportKpis');
    const weeklyReportText = document.getElementById('weeklyReportText');
    const weeklyReportNotice = document.getElementById('weeklyReportNotice');
    const copyWeeklyReportBtn = document.getElementById('copyWeeklyReportBtn');
    const downloadWeeklyReportBtn = document.getElementById('downloadWeeklyReportBtn');
    let weeklyReportDownloadUrl = '';
    let dashboardLoading = false;
    let dashboardReloadQueued = false;
    const DASHBOARD_REFRESH_MS = 60000;
    const THEME_AUTO_REFRESH_MS = 60000;
    const THEME_KEY = 'asset-dashboard-theme-mode';
    const THEME_MODES = ['auto', 'light', 'dark'];

    const esc = (value) => String(value ?? '').replace(/[&<>"']/g, ch => ({
      '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'
    }[ch]));

    function preferredThemeForTime() {
      const hour = new Date().getHours();
      return hour >= 7 && hour < 18 ? 'light' : 'dark';
    }

    function currentThemeMode() {
      let saved = 'auto';
      try {
        saved = localStorage.getItem(THEME_KEY) || 'auto';
      } catch (_) {}
      return THEME_MODES.includes(saved) ? saved : 'auto';
    }

    function applyTheme() {
      const mode = currentThemeMode();
      const actualTheme = mode === 'auto' ? preferredThemeForTime() : mode;
      document.body.dataset.theme = actualTheme;
      themeToggleBtn.textContent = mode === 'auto'
        ? `自动外观：${actualTheme === 'dark' ? '夜晚' : '白天'}`
        : actualTheme === 'dark' ? '夜晚外观' : '白天外观';
      themeToggleBtn.title = '点击切换：自动外观 / 白天外观 / 夜晚外观';
    }

    function toggleThemeMode() {
      const mode = currentThemeMode();
      const nextMode = THEME_MODES[(THEME_MODES.indexOf(mode) + 1) % THEME_MODES.length];
      try {
        localStorage.setItem(THEME_KEY, nextMode);
      } catch (_) {}
      applyTheme();
    }

    function cleanQuery(value) {
      return String(value || '').replace(/\s+/g, '').replace(/\u3000/g, '').trim();
    }

    function isComposingInput(input) {
      return input.dataset.composing === '1';
    }

    function isComposingEvent(event) {
      return Boolean(event.isComposing || event.keyCode === 229 || isComposingInput(event.target));
    }

    function bindCompositionGuard(input, onCommit) {
      input.addEventListener('compositionstart', () => {
        input.dataset.composing = '1';
      });
      input.addEventListener('compositionend', () => {
        input.dataset.composing = '0';
        if (onCommit) onCommit();
      });
    }

    async function getJson(url) {
      const res = await fetch(url, { cache: 'no-store' });
      if (!res.ok) throw new Error(`HTTP ${res.status}`);
      return await res.json();
    }

    async function postJson(url, payload) {
      const res = await fetch(url, {
        method: 'POST',
        headers: { 'Content-Type': 'application/json' },
        cache: 'no-store',
        body: JSON.stringify(payload || {}),
      });
      const data = await res.json().catch(() => ({}));
      if (!res.ok || data.ok === false) throw new Error(data.message || `HTTP ${res.status}`);
      return data;
    }

    function tone(status) {
      if (['入库', '维修完成已重新入库', '已换新'].includes(status)) return 'tone-green';
      if (status === '领用') return 'tone-blue';
      if (status === '借用中') return 'tone-purple';
      if (status === '归还') return 'tone-soft';
      if (status === '退租') return 'tone-orange';
      if (status === '维修中') return 'tone-yellow';
      if (status === '删除资产') return 'tone-gray';
      return 'tone-soft';
    }

    function badge(status) {
      return `<span class="badge ${tone(status)}">${esc(status || '-')}</span>`;
    }

    function kpi(label, value) {
      const displayValue = value === null || value === undefined || value === '' ? '-' : value;
      return `<div class="kpi"><div class="label">${esc(label)}</div><div class="value">${displayValue}</div></div>`;
    }

    function mini(label, value) {
      const displayValue = value === null || value === undefined || value === '' ? '-' : value;
      return `<div class="mini-card"><div class="mini-label">${esc(label)}</div><div class="value">${displayValue}</div></div>`;
    }

    function chips(items, labelKey, emptyText) {
      if (!items || !items.length) return `<div class="muted">${esc(emptyText || '暂无数据')}</div>`;
      return items.map(item => `
        <div class="chip">
          <span>${esc(item[labelKey] || '-')}</span>
          <strong>${esc(item.count)}</strong>
        </div>
      `).join('');
    }

    async function loadDashboard() {
      if (dashboardLoading) {
        dashboardReloadQueued = true;
        return;
      }
      dashboardLoading = true;
      try {
        const data = await getJson('/api/dashboard');
        const status = data.status || {};
        const overview = data.overview || {};
        const currentAssetCount = overview.currentAssetCount ?? '-';
        const historicalAssetCount = overview.historicalAssetCount ?? overview.assetCount ?? 0;
        metaEl.textContent = `当前在册 ${currentAssetCount} 台 / 历史追踪 ${historicalAssetCount} 台 / 记录 ${status.rowCount || 0} 条 / 数据时间 ${status.sourceMtime || '-'}`;
        overviewEl.innerHTML = [
          kpi('当前在册资产', currentAssetCount),
          kpi('历史追踪资产', historicalAssetCount),
          kpi('本周操作', overview.weeklyCount ?? 0),
          kpi('统计周期', `${esc(overview.weekStart || '-')} 至 ${esc(overview.weekEnd || '-')}`),
        ].join('');
        weeklyStatusEl.innerHTML = chips(data.weeklyStatusCounts, 'status', '本周暂无状态结果');
        weeklyActionEl.innerHTML = chips(data.weeklyActionCounts, 'action', '本周暂无操作类型');
        weeklyCategoryEl.innerHTML = chips((data.weeklyReport || {}).categoryCounts, 'category', '本周暂无资产分类数据');
        renderWeeklyReport(data.weeklyReport || {});
        renderRecent(data.recent || []);
      } catch (err) {
        metaEl.textContent = `数据读取失败：${err.message}`;
      } finally {
        dashboardLoading = false;
        if (dashboardReloadQueued) {
          dashboardReloadQueued = false;
          loadDashboard();
        }
      }
    }

    function renderWeeklyReport(report) {
      const delta = Number(report.delta || 0);
      const deltaText = delta > 0 ? `+${delta}` : String(delta);
      weeklyReportKpis.innerHTML = [
        mini('周报周期', `${esc(report.weekLabel || '-')} / ${esc(report.weekStart || '-')} 至 ${esc(report.weekEnd || '-')}`),
        mini('本周操作', esc(report.operationCount || 0)),
        mini('涉及资产', esc(report.assetCount || 0)),
        mini('较上周同期', esc(deltaText)),
      ].join('');
      weeklyReportText.value = report.copyText || '暂无周报内容。';
      const download = report.download || {};
      weeklyReportDownloadUrl = download.exists ? download.url : '';
      downloadWeeklyReportBtn.disabled = !weeklyReportDownloadUrl;
      weeklyReportNotice.textContent = download.exists
        ? `Excel归档：${download.fileName || '资产周度汇总.xlsx'}，更新时间 ${download.modifiedAt || '-'}`
        : '当前没有可下载的周报 Excel，网页周报文字仍可直接复制。';
    }

    async function searchAsset() {
      if (isComposingInput(assetQuery)) return;
      const q = cleanQuery(assetQuery.value);
      assetQuery.value = q;
      if (!q) {
        assetNotice.textContent = '请输入资产编号、序列号或资产名称关键字。';
        assetResult.classList.add('hidden');
        return;
      }
      assetNotice.textContent = '正在查询...';
      try {
        const data = await getJson(`/api/search?q=${encodeURIComponent(q)}`);
        if (data.ok) {
          assetNotice.textContent = '已找到资产。';
          renderAsset(data.asset);
          renderSuggestions([]);
        } else {
          assetNotice.textContent = data.message || '未找到。';
          assetResult.classList.add('hidden');
          renderSuggestions(data.suggestions || []);
        }
      } catch (err) {
        assetNotice.textContent = `查询失败：${err.message}`;
      }
    }

    async function searchPerson() {
      if (isComposingInput(personQuery)) return;
      const q = cleanQuery(personQuery.value);
      personQuery.value = q;
      if (!q) {
        personNotice.textContent = '请输入员工姓名、部门或使用对象关键字。';
        personResult.classList.add('hidden');
        return;
      }
      personNotice.textContent = '正在查询...';
      try {
        const data = await getJson(`/api/person?q=${encodeURIComponent(q)}`);
        if (!data.ok) {
          personNotice.textContent = data.message || '没有找到。';
        } else {
          personNotice.textContent = `当前领用 ${data.currentCount || 0} 台，相关历史 ${data.activityCount || 0} 条。`;
        }
        renderPerson(data);
      } catch (err) {
        personNotice.textContent = `查询失败：${err.message}`;
      }
    }

    async function loadSuggestions() {
      if (isComposingInput(assetQuery)) return;
      const q = cleanQuery(assetQuery.value);
      if (!q) {
        renderSuggestions([]);
        return;
      }
      try {
        const data = await getJson(`/api/suggest?q=${encodeURIComponent(q)}`);
        renderSuggestions(data.suggestions || []);
      } catch (_) {}
    }

    function renderAsset(asset) {
      assetResult.classList.remove('hidden');
      assetSummary.innerHTML = [
        mini('资产编号', esc(asset.assetTag)),
        mini('当前状态', badge(asset.currentStatus)),
        mini('当前使用人', esc(asset.currentHolder || '-')),
        mini('最近操作', esc(asset.latestTime || '-')),
        mini('资产名称', esc(asset.assetName || '-')),
        mini('资产分类', esc(asset.category || '-')),
        mini('序列号', esc(asset.serial || '-')),
        mini('总操作次数', esc(asset.totalCount ?? 0)),
      ].join('');
      assetIssues.innerHTML = asset.issues && asset.issues.length
        ? asset.issues.map(item => `<span class="issue">${esc(item)}</span>`).join('')
        : '<span class="issue issue-ok">当前未发现明显提醒</span>';
      assetHistory.innerHTML = (asset.history || []).map(row => `
        <tr>
          <td>${esc(row.actionTime)}</td>
          <td>${badge(row.status)}</td>
          <td>${esc(row.action)}</td>
          <td>${esc(row.target)}</td>
          <td>${esc(row.operator)}</td>
          <td>${esc(row.note)}</td>
        </tr>
      `).join('') || '<tr><td colspan="6" class="muted">没有历史记录。</td></tr>';
      assetResult.scrollIntoView({ behavior: 'smooth', block: 'start' });
    }

    function renderPerson(data) {
      personResult.classList.remove('hidden');
      personSummary.innerHTML = [
        mini('查询关键字', esc(data.query || '-')),
        mini('当前领用资产', esc(data.currentCount || 0)),
        mini('相关历史记录', esc(data.activityCount || 0)),
        mini('结果状态', data.ok ? '<span class="badge tone-green">已匹配</span>' : '<span class="badge tone-yellow">未匹配</span>'),
      ].join('');
      personHoldings.innerHTML = (data.holdings || []).map(item => `
        <tr class="clickable-row" data-asset="${esc(item.assetTag)}">
          <td>${esc(item.assetTag)}</td>
          <td>${badge(item.currentStatus)}</td>
          <td>${esc(item.assetName)}</td>
          <td>${esc(item.category)}</td>
          <td>${esc(item.latestTime)}</td>
        </tr>
      `).join('') || '<tr><td colspan="5" class="muted">没有当前领用资产。</td></tr>';
      personHistory.innerHTML = (data.history || []).map(row => `
        <tr>
          <td>${esc(row.assetTag)}</td>
          <td>${esc(row.actionTime)}</td>
          <td>${badge(row.status)}</td>
          <td>${esc(row.target || row.operator)}</td>
        </tr>
      `).join('') || '<tr><td colspan="4" class="muted">没有相关历史记录。</td></tr>';
      personHoldings.querySelectorAll('[data-asset]').forEach(row => {
        row.addEventListener('click', () => {
          assetQuery.value = row.dataset.asset;
          searchAsset();
        });
      });
      personResult.scrollIntoView({ behavior: 'smooth', block: 'start' });
    }

    function renderSuggestions(items) {
      if (!items.length) {
        suggestionsEl.classList.add('hidden');
        suggestionsEl.innerHTML = '';
        return;
      }
      suggestionsEl.classList.remove('hidden');
      suggestionsEl.innerHTML = items.map(item => `
        <div class="asset-card" data-asset="${esc(item.assetTag)}">
          <strong>${esc(item.assetTag)}</strong>
          <span>${esc(item.assetName || '-')}</span>
          <span>${esc(item.category || '-')} / ${esc(item.serial || '-')}</span>
          <span>${badge(item.currentStatus)} ${item.currentHolder ? ' / ' + esc(item.currentHolder) : ''}</span>
        </div>
      `).join('');
      suggestionsEl.querySelectorAll('.asset-card').forEach(card => {
        card.addEventListener('click', () => {
          assetQuery.value = card.dataset.asset;
          searchAsset();
        });
      });
    }

    function renderRecent(items) {
      recentEl.innerHTML = items.map(row => `
        <tr>
          <td>${esc(row.assetTag)}</td>
          <td>${esc(row.actionTime)}</td>
          <td>${badge(row.status)}</td>
          <td>${esc(row.target || row.operator || '-')}</td>
        </tr>
      `).join('') || '<tr><td colspan="4" class="muted">暂无最近操作。</td></tr>';
    }

    let suggestTimer = null;
    bindCompositionGuard(assetQuery, () => {
      clearTimeout(suggestTimer);
      suggestTimer = setTimeout(loadSuggestions, 120);
    });
    bindCompositionGuard(personQuery);
    assetQuery.addEventListener('input', () => {
      if (isComposingInput(assetQuery)) return;
      clearTimeout(suggestTimer);
      suggestTimer = setTimeout(loadSuggestions, 180);
    });
    assetQuery.addEventListener('keydown', event => {
      if (isComposingEvent(event)) return;
      if (event.key === 'Enter') searchAsset();
    });
    personQuery.addEventListener('keydown', event => {
      if (isComposingEvent(event)) return;
      if (event.key === 'Enter') searchPerson();
    });
    document.getElementById('assetSearchBtn').addEventListener('click', searchAsset);
    document.getElementById('assetClearBtn').addEventListener('click', () => {
      assetQuery.value = '';
      assetNotice.textContent = '复制编号里带换行或空格也可以自动处理。';
      assetResult.classList.add('hidden');
      renderSuggestions([]);
      assetQuery.focus();
    });
    document.getElementById('personSearchBtn').addEventListener('click', searchPerson);
    document.getElementById('personClearBtn').addEventListener('click', () => {
      personQuery.value = '';
      personNotice.textContent = '可查看当前领用资产和相关历史记录。';
      personResult.classList.add('hidden');
      personQuery.focus();
    });
    copyWeeklyReportBtn.addEventListener('click', async () => {
      weeklyReportText.select();
      try {
        await navigator.clipboard.writeText(weeklyReportText.value);
        weeklyReportNotice.textContent = '周报文字已复制，可以直接粘贴到周报里。';
      } catch (_) {
        document.execCommand('copy');
        weeklyReportNotice.textContent = '周报文字已复制，可以直接粘贴到周报里。';
      }
    });
    downloadWeeklyReportBtn.addEventListener('click', () => {
      if (weeklyReportDownloadUrl) window.location.href = weeklyReportDownloadUrl;
    });
    themeToggleBtn.addEventListener('click', toggleThemeMode);

    applyTheme();
    loadDashboard();
    setInterval(() => {
      if (currentThemeMode() === 'auto') applyTheme();
    }, THEME_AUTO_REFRESH_MS);
    setInterval(loadDashboard, DASHBOARD_REFRESH_MS);
    assetQuery.focus();
  </script>
</body>
</html>
"""


class Handler(BaseHTTPRequestHandler):
    store: LifecycleStore

    def do_HEAD(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            return
        if parsed.path == "/download/weekly-report":
            report_path = generate_weekly_report_file()
            if not report_path:
                self.send_error(HTTPStatus.NOT_FOUND, "Weekly report not found")
                return
            self.send_file(report_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", head_only=True)
            return
        self.send_error(HTTPStatus.NOT_FOUND, "Not found")

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self.send_text(HTML, "text/html; charset=utf-8")
            return
        if parsed.path == "/api/dashboard":
            self.send_json(self.store.dashboard())
            return
        if parsed.path == "/api/weekly-report":
            self.send_json(self.store.weekly_report())
            return
        if parsed.path == "/api/status":
            self.send_json(self.store.status())
            return
        if parsed.path == "/api/search":
            query = parse_qs(parsed.query).get("q", [""])[0]
            self.send_json(self.store.search(query))
            return
        if parsed.path == "/api/person":
            query = parse_qs(parsed.query).get("q", [""])[0]
            self.send_json(self.store.person_search(query))
            return
        if parsed.path == "/api/suggest":
            query = parse_qs(parsed.query).get("q", [""])[0]
            self.send_json({"suggestions": self.store.suggestions(query)})
            return
        if parsed.path == "/download/weekly-report":
            report_path = generate_weekly_report_file()
            if not report_path:
                self.send_error(HTTPStatus.NOT_FOUND, "Weekly report not found")
                return
            self.send_file(report_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            return
        self.send_error(HTTPStatus.NOT_FOUND, "Not found")

    def do_POST(self) -> None:
        self.send_error(HTTPStatus.NOT_FOUND, "Not found")

    def log_message(self, format: str, *args: Any) -> None:
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        sys.stdout.write(f"[{stamp}] {self.address_string()} {format % args}\n")
        sys.stdout.flush()

    def send_json(self, payload: Any) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json_body(self) -> dict[str, Any]:
        length = int(self.headers.get("Content-Length", "0") or "0")
        if length <= 0:
            return {}
        raw = self.rfile.read(length)
        try:
            data = json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            return {}
        return data if isinstance(data, dict) else {}

    def send_text(self, text: str, content_type: str) -> None:
        body = text.encode("utf-8")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_file(self, path: Path, content_type: str, head_only: bool = False) -> None:
        filename = path.name
        encoded_filename = quote(filename)
        size = path.stat().st_size
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(size))
        self.send_header(
            "Content-Disposition",
            f"attachment; filename*=UTF-8''{encoded_filename}",
        )
        self.end_headers()
        if head_only:
            return
        with path.open("rb") as handle:
            while True:
                chunk = handle.read(1024 * 1024)
                if not chunk:
                    break
                self.wfile.write(chunk)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run asset lifecycle web app.")
    parser.add_argument("--source", type=Path, default=SOURCE_PATH)
    parser.add_argument("--host", default=HOST)
    parser.add_argument("--port", type=int, default=PORT)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    Handler.store = LifecycleStore(args.source)
    server = ThreadingHTTPServer((args.host, args.port), Handler)
    print(f"Asset lifecycle web app listening on http://{args.host}:{args.port}/")
    print(f"Source workbook: {args.source}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("Stopping.")
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
