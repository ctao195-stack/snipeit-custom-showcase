#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate historical month-end Snipe-IT asset fee workbooks from the database.

This script is for backfilling missed month-end reports. It reconstructs the
asset scope using created_at, deleted_at, and retire/restore action logs, then
writes XLSX files without touching 最新资产费用清单.xlsx.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import tempfile
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


AUTOMATION_ROOT = Path(os.getenv("SNIPEIT_AUTOMATION_ROOT", r"D:\SnipeIT-Automation"))
DEFAULT_OUTPUT_DIR = Path(
    os.getenv("SNIPEIT_MONTHLY_OUTPUT_DIR", str(AUTOMATION_ROOT / "output" / "monthly-fees"))
)
DEFAULT_LOG_PATH = Path(
    os.getenv(
        "SNIPEIT_HISTORICAL_REPORT_LOG",
        str(AUTOMATION_ROOT / "logs" / "historical-monthly-fee-report.log"),
    )
)
PHP_EXE = Path(os.getenv("SNIPEIT_PHP_EXE", r"C:\php\php.exe"))
SNIPEIT_ENV_PATH = Path(
    os.getenv("SNIPEIT_APP_ENV_PATH", r"C:\inetpub\snipe-it\.env")
)
SNIPEIT_BASE_URL = os.getenv("SNIPEIT_URL", "https://snipe-it.example.com").rstrip("/")
RENTAL_PREFIX = os.getenv("SNIPEIT_RENTAL_PREFIX", "RENT-").upper()
OWNED_PREFIX = os.getenv("SNIPEIT_OWNED_PREFIX", "OWN-").upper()

FONT_NAME = "Microsoft YaHei"
PRIMARY = "1F4E78"
PRIMARY_DARK = "17365D"
PALE_BLUE = "F6FAFD"
BORDER_GRAY = "D9E2EC"
WARNING = "FFF2CC"

DETAIL_HEADERS = [
    "序号",
    "资产编号",
    "资产名称",
    "资产来源",
    "分类",
    "型号",
    "型号编码",
    "序列号",
    "月末是否在账",
    "当前状态",
    "使用人/归属",
    "部门",
    "公司",
    "位置",
    "供应商",
    "采购日期",
    "订单号",
    "金额",
    "资产创建时间",
    "资产删除时间",
    "月底前最后退租/恢复动作",
    "资产链接",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Backfill historical Snipe-IT month-end fee reports.")
    parser.add_argument("--dates", nargs="+", required=True, help="Month-end dates, e.g. 2026-04-30 2026-05-31")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--log", type=Path, default=DEFAULT_LOG_PATH)
    parser.add_argument("--suffix", default="", help="Optional filename suffix.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    for value in args.dates:
        cutoff = parse_cutoff(value)
        payload = fetch_historical_assets(cutoff)
        rows = [normalize_row(row) for row in payload["rows"]]
        rows.sort(key=lambda item: natural_sort_key(item["资产编号"]))

        month_label = cutoff.strftime("%Y-%m")
        output_path = args.output_dir / f"{month_label}资产费用清单{args.suffix}.xlsx"
        create_workbook(rows, output_path, cutoff, payload)

        total = sum((row["金额"] for row in rows), Decimal("0"))
        message = f"generated historical: {output_path}; cutoff={cutoff:%Y-%m-%d}; assets={len(rows)}; total={format_decimal(total)}"
        write_log(args.log, message)
        print(message)

    return 0


def parse_cutoff(value: str) -> datetime:
    return datetime.strptime(value + " 23:59:59", "%Y-%m-%d %H:%M:%S")


def fetch_historical_assets(cutoff: datetime) -> dict[str, Any]:
    php_code = build_php_query(cutoff.strftime("%Y-%m-%d %H:%M:%S"))
    with tempfile.NamedTemporaryFile("w", suffix=".php", delete=False, encoding="utf-8-sig") as handle:
        handle.write(php_code)
        temp_path = Path(handle.name)

    try:
        result = subprocess.run(
            [str(PHP_EXE), str(temp_path)],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
    finally:
        temp_path.unlink(missing_ok=True)

    if result.returncode != 0:
        raise RuntimeError(f"PHP query failed: {result.stderr.strip() or result.stdout.strip()}")

    return json.loads(result.stdout.lstrip("\ufeff"))


def build_php_query(cutoff: str) -> str:
    env_path = json.dumps(str(SNIPEIT_ENV_PATH))

    return f"""<?php
$envPath = {env_path};
$env = [];
foreach (file($envPath, FILE_IGNORE_NEW_LINES | FILE_SKIP_EMPTY_LINES) as $line) {{
    $line = trim($line);
    if ($line === '' || $line[0] === '#' || strpos($line, '=') === false) {{
        continue;
    }}
    [$key, $value] = explode('=', $line, 2);
    $env[trim($key)] = trim($value, " \\t\\n\\r\\0\\x0B\\"'");
}}

$pdo = new PDO(
    "mysql:host={{$env['DB_HOST']}};port={{$env['DB_PORT']}};dbname={{$env['DB_DATABASE']}};charset=utf8mb4",
    $env['DB_USERNAME'],
    $env['DB_PASSWORD'],
    [PDO::ATTR_ERRMODE => PDO::ERRMODE_EXCEPTION, PDO::ATTR_DEFAULT_FETCH_MODE => PDO::FETCH_ASSOC]
);

$cutoff = '{cutoff}';
$assetModel = 'App\\\\\\\\Models\\\\\\\\Asset';

$sql = "
select
    a.id,
    a.asset_tag,
    a.name as asset_name,
    a.serial,
    a.purchase_cost,
    a.purchase_date,
    a.order_number,
    a.created_at,
    a.deleted_at,
    sl.name as current_status,
    mdl.name as model_name,
    mdl.model_number,
    cat.name as category_name,
    mf.name as manufacturer_name,
    sup.name as supplier_name,
    co.name as company_name,
    loc.name as location_name,
    rtd.name as default_location_name,
    u.first_name,
    u.last_name,
    u.username,
    dep.name as department_name,
    terminal.action_type as terminal_action,
    terminal.created_at as terminal_action_at
from assets a
left join status_labels sl on sl.id = a.status_id
left join models mdl on mdl.id = a.model_id
left join categories cat on cat.id = mdl.category_id
left join manufacturers mf on mf.id = mdl.manufacturer_id
left join suppliers sup on sup.id = a.supplier_id
left join companies co on co.id = a.company_id
left join locations loc on loc.id = a.location_id
left join locations rtd on rtd.id = a.rtd_location_id
left join users u on a.assigned_type = 'App\\\\\\\\Models\\\\\\\\User' and u.id = a.assigned_to
left join departments dep on dep.id = u.department_id
left join (
    select al.item_id, al.action_type, al.created_at
    from action_logs al
    inner join (
        select item_id, max(id) as max_id
        from action_logs
        where item_type = :asset_model
          and action_type in ('retire', 'restore')
          and created_at <= :cutoff
        group by item_id
    ) latest on latest.max_id = al.id
) terminal on terminal.item_id = a.id
where a.created_at <= :cutoff
  and (a.deleted_at is null or a.deleted_at > :cutoff)
  and (terminal.action_type is null or terminal.action_type <> 'retire')
  and coalesce(sl.name, '') <> '退租'
order by a.asset_tag
";

$stmt = $pdo->prepare($sql);
$stmt->execute(['asset_model' => $assetModel, 'cutoff' => $cutoff]);
$rows = $stmt->fetchAll();

$excludedSql = "
select al.action_type, count(*) as c
from action_logs al
where al.item_type = :asset_model
  and al.created_at <= :cutoff
  and al.action_type in ('retire', 'delete', 'restore')
group by al.action_type
order by al.action_type
";
$excludedStmt = $pdo->prepare($excludedSql);
$excludedStmt->execute(['asset_model' => $assetModel, 'cutoff' => $cutoff]);

echo json_encode([
    'cutoff' => $cutoff,
    'rows' => $rows,
    'action_counts' => $excludedStmt->fetchAll(),
], JSON_UNESCAPED_UNICODE | JSON_UNESCAPED_SLASHES);
?>"""


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    asset_tag = text(row.get("asset_tag"))
    return {
        "序号": 0,
        "资产编号": asset_tag,
        "资产名称": text(row.get("asset_name")),
        "资产来源": classify_ownership(asset_tag),
        "分类": text(row.get("category_name")),
        "型号": text(row.get("model_name")),
        "型号编码": text(row.get("model_number")),
        "序列号": text(row.get("serial")),
        "月末是否在账": "是",
        "当前状态": text(row.get("current_status")),
        "使用人/归属": user_name(row),
        "部门": text(row.get("department_name")),
        "公司": text(row.get("company_name")),
        "位置": text(row.get("location_name")) or text(row.get("default_location_name")),
        "供应商": text(row.get("supplier_name")),
        "采购日期": text(row.get("purchase_date")),
        "订单号": text(row.get("order_number")),
        "金额": money(row.get("purchase_cost")),
        "资产创建时间": text(row.get("created_at")),
        "资产删除时间": text(row.get("deleted_at")),
        "月底前最后退租/恢复动作": terminal_text(row),
        "资产链接": f"{SNIPEIT_BASE_URL}/hardware/{row.get('id')}" if row.get("id") else "",
    }


def user_name(row: dict[str, Any]) -> str:
    full = " ".join(part for part in [text(row.get("last_name")), text(row.get("first_name"))] if part).strip()
    return full or text(row.get("username"))


def terminal_text(row: dict[str, Any]) -> str:
    action = text(row.get("terminal_action"))
    when = text(row.get("terminal_action_at"))
    return f"{action} {when}".strip()


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def money(value: Any) -> Decimal:
    raw = text(value)
    if not raw:
        return Decimal("0")
    cleaned = re.sub(r"[^0-9.\-]", "", raw.replace(",", ""))
    if cleaned in {"", "-", ".", "-."}:
        return Decimal("0")
    try:
        return Decimal(cleaned).quantize(Decimal("0.01"))
    except InvalidOperation:
        return Decimal("0")


def classify_ownership(asset_tag: str) -> str:
    normalized = asset_tag.upper()
    if normalized.startswith(RENTAL_PREFIX):
        return "租赁"
    if normalized.startswith(OWNED_PREFIX):
        return "自购"
    return "未识别"


def natural_sort_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value or "")]


def create_workbook(rows: list[dict[str, Any]], output_path: Path, cutoff: datetime, payload: dict[str, Any]) -> None:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "费用清单"
    build_detail_sheet(detail, rows)

    summary = workbook.create_sheet("费用汇总")
    build_summary_sheet(summary, rows, cutoff)

    notes = workbook.create_sheet("历史口径说明")
    build_notes_sheet(notes, rows, cutoff, payload)

    workbook.save(output_path)


def build_detail_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    sheet.append(DETAIL_HEADERS)
    for index, row in enumerate(rows, start=1):
        row["序号"] = index
        sheet.append([row[header] for header in DETAIL_HEADERS])

    apply_sheet_style(sheet, len(rows) + 1, len(DETAIL_HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    amount_col = DETAIL_HEADERS.index("金额") + 1
    link_col = DETAIL_HEADERS.index("资产链接") + 1
    for row_number in range(2, len(rows) + 2):
        sheet.cell(row_number, amount_col).number_format = '#,##0.00'
        link_cell = sheet.cell(row_number, link_col)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"

    add_table(sheet, "HistoricalMonthlyFeeDetail")
    set_widths(
        sheet,
        {
            "A": 8,
            "B": 24,
            "C": 18,
            "D": 10,
            "E": 16,
            "F": 26,
            "G": 18,
            "H": 24,
            "I": 14,
            "J": 14,
            "K": 20,
            "L": 16,
            "M": 16,
            "N": 16,
            "O": 18,
            "P": 14,
            "Q": 22,
            "R": 12,
            "S": 22,
            "T": 22,
            "U": 24,
            "V": 42,
        },
    )


def build_summary_sheet(sheet: Any, rows: list[dict[str, Any]], cutoff: datetime) -> None:
    total_cost = sum((row["金额"] for row in rows), Decimal("0"))
    zero_cost = sum(1 for row in rows if row["金额"] == 0)

    sheet["A1"] = f"{cutoff.strftime('%Y-%m')} 历史月末资产费用汇总"
    sheet["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=PRIMARY)
    sheet.merge_cells("A1:D1")

    summary_rows = [
        ("统计截至", cutoff.strftime("%Y-%m-%d 23:59:59")),
        ("资产总数", len(rows)),
        ("费用合计", float(total_cost)),
        ("金额为空/0 的资产", zero_cost),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for index, item in enumerate(summary_rows, start=3):
        sheet.cell(index, 1).value = item[0]
        sheet.cell(index, 2).value = item[1]
        sheet.cell(index, 1).font = Font(name=FONT_NAME, bold=True, color=PRIMARY_DARK)
        if item[0] == "费用合计":
            sheet.cell(index, 2).number_format = '#,##0.00'

    row_cursor = 10
    row_cursor = add_group_summary(sheet, row_cursor, "按资产来源", rows, "资产来源")
    row_cursor = add_group_summary(sheet, row_cursor + 2, "按分类", rows, "分类")
    row_cursor = add_group_summary(sheet, row_cursor + 2, "按当前状态", rows, "当前状态")
    add_group_summary(sheet, row_cursor + 2, "按公司", rows, "公司")

    set_widths(sheet, {"A": 28, "B": 16, "C": 16, "D": 18})
    sheet.freeze_panes = "A3"


def add_group_summary(sheet: Any, start_row: int, title: str, rows: list[dict[str, Any]], key: str) -> int:
    grouped: dict[str, dict[str, Decimal | int]] = defaultdict(lambda: {"count": 0, "amount": Decimal("0")})
    for row in rows:
        label = row.get(key) or "未填写"
        grouped[str(label)]["count"] += 1
        grouped[str(label)]["amount"] += row["金额"]

    sheet.cell(start_row, 1).value = title
    sheet.cell(start_row, 1).font = Font(name=FONT_NAME, size=12, bold=True, color=PRIMARY_DARK)
    for col, title_value in enumerate([key, "资产数", "金额合计", "占比"], start=1):
        cell = sheet.cell(start_row + 1, col)
        cell.value = title_value
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY)
        cell.alignment = Alignment(horizontal="center")

    total_amount = sum((row["金额"] for row in rows), Decimal("0"))
    current = start_row + 2
    for label, item in sorted(grouped.items(), key=lambda entry: (-entry[1]["amount"], entry[0])):
        amount = item["amount"]
        ratio = float(amount / total_amount) if total_amount else 0
        sheet.cell(current, 1).value = label
        sheet.cell(current, 2).value = item["count"]
        sheet.cell(current, 3).value = float(amount)
        sheet.cell(current, 3).number_format = '#,##0.00'
        sheet.cell(current, 4).value = ratio
        sheet.cell(current, 4).number_format = "0.00%"
        current += 1

    apply_sheet_style(sheet, current - 1, 4, start_row=start_row + 1)
    return current


def build_notes_sheet(sheet: Any, rows: list[dict[str, Any]], cutoff: datetime, payload: dict[str, Any]) -> None:
    action_counts = ", ".join(f"{item['action_type']}={item['c']}" for item in payload.get("action_counts", []))
    lines = [
        ("报表类型", "历史月末补生成"),
        ("统计截至", cutoff.strftime("%Y-%m-%d 23:59:59")),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("历史在账口径", "资产创建时间 <= 月底；删除时间为空或晚于月底；月底前最后一次 retire/restore 动作不是 retire；当前状态不是退租。"),
        ("费用口径", "金额取当前资产主数据 purchase_cost。若历史月份采购价格曾被后续修改，需要以当时备份核对。"),
        ("状态/使用人口径", "当前状态和使用人来自当前资产主数据，主要用于识别，不作为历史月末领用关系的唯一依据。"),
        ("动作统计", action_counts or "无"),
        ("提醒", "本报表用于补 2026 年 4-6 月漏统费用；如需财务级严格追溯，可再对比当月数据库备份。"),
    ]
    for row_number, (key, value) in enumerate(lines, start=1):
        sheet.cell(row_number, 1).value = key
        sheet.cell(row_number, 2).value = value
    apply_sheet_style(sheet, len(lines), 2)
    set_widths(sheet, {"A": 18, "B": 120})

    if any(row["金额"] == 0 for row in rows):
        row_number = len(lines) + 2
        sheet.cell(row_number, 1).value = "金额提醒"
        sheet.cell(row_number, 2).value = "存在金额为空或 0 的资产，请确认 Snipe-IT 中是否已维护采购价格/费用。"
        sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor=WARNING)
        sheet.cell(row_number, 2).fill = PatternFill("solid", fgColor=WARNING)


def apply_sheet_style(sheet: Any, max_row: int, max_col: int, start_row: int = 1) -> None:
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(min_row=start_row, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=10, color="1F2933")
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for cell in sheet[start_row]:
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number in range(start_row + 1, max_row + 1):
        fill = PatternFill("solid", fgColor=PALE_BLUE if row_number % 2 == 0 else "FFFFFF")
        for col in range(1, max_col + 1):
            sheet.cell(row_number, col).fill = fill


def add_table(sheet: Any, name: str) -> None:
    if sheet.max_row < 2:
        return
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def set_widths(sheet: Any, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows():
        sheet.row_dimensions[row[0].row].height = 22


def format_decimal(value: Decimal) -> str:
    return f"{value:,.2f}"


def write_log(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {message}"
    with path.open("a", encoding="utf-8") as handle:
        handle.write(line + os.linesep)


if __name__ == "__main__":
    raise SystemExit(main())
