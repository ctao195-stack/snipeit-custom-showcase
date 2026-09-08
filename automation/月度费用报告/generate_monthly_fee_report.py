#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate a monthly Snipe-IT asset fee workbook.

The script is read-only against Snipe-IT. It uses the existing API token from
the configured environment file and writes a workbook to the output directory.
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_ENV_PATH = Path(r"D:\snipe-it-excel-sync\.env.snipeit-excel")
DEFAULT_OUTPUT_DIR = Path(r"D:\月度费用报告")
DEFAULT_LOG_PATH = Path(r"D:\snipe-it-excel-sync\月度费用报告\月度费用报告.log")

FONT_NAME = "Microsoft YaHei"
PRIMARY = "1F4E78"
PRIMARY_DARK = "17365D"
PALE_BLUE = "F6FAFD"
BORDER_GRAY = "D9E2EC"
WARNING = "FFF2CC"

DETAIL_HEADERS = [
    "序号",
    "资产编号",
    "资产名称",
    "资产来源",
    "分类",
    "型号",
    "型号编码",
    "序列号",
    "当前状态",
    "使用人/归属",
    "部门",
    "公司",
    "位置",
    "供应商",
    "采购日期",
    "订单号",
    "金额",
    "资产链接",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate monthly Snipe-IT asset fee report.")
    parser.add_argument("--env", type=Path, default=DEFAULT_ENV_PATH)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--log", type=Path, default=DEFAULT_LOG_PATH)
    parser.add_argument("--run-date", help="YYYY-MM-DD. Defaults to today.")
    parser.add_argument("--force", action="store_true", help="Generate even when today is not month end.")
    parser.add_argument("--page-size", type=int, default=500)
    parser.add_argument("--max-pages", type=int, default=100)
    parser.add_argument("--suffix", default="", help="Optional filename suffix, such as _test.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    run_date = parse_run_date(args.run_date)

    if not args.force and not is_last_day(run_date):
        message = f"skip: {run_date.isoformat()} is not the last day of the month"
        write_log(args.log, message)
        print(message)
        return 0

    env = read_env(args.env)
    base_url = required(env, "SNIPEIT_URL").rstrip("/") + "/"
    api_token = required(env, "SNIPEIT_API_TOKEN")
    verify_ssl = env.get("SNIPEIT_VERIFY_SSL", "true").strip().lower() not in {"0", "false", "no"}
    rental_prefix = env.get("SNIPEIT_RENTAL_PREFIX", "RENT-").strip().upper()
    owned_prefix = env.get("SNIPEIT_OWNED_PREFIX", "OWN-").strip().upper()

    session = requests.Session()
    session.headers.update(
        {
            "Accept": "application/json",
            "Authorization": f"Bearer {api_token}",
        }
    )

    rows = fetch_assets(
        session=session,
        base_url=base_url,
        verify_ssl=verify_ssl,
        page_size=max(50, args.page_size),
        max_pages=max(1, args.max_pages),
    )
    report_rows = [normalize_asset(row, base_url, rental_prefix, owned_prefix) for row in rows]
    report_rows.sort(key=lambda item: natural_sort_key(item["资产编号"]))

    args.output_dir.mkdir(parents=True, exist_ok=True)
    month_label = run_date.strftime("%Y-%m")
    output_path = args.output_dir / f"{month_label}资产费用清单{args.suffix}.xlsx"
    latest_path = args.output_dir / "最新资产费用清单.xlsx"

    create_workbook(report_rows, output_path, base_url, run_date, rental_prefix, owned_prefix)
    shutil.copyfile(output_path, latest_path)

    total_cost = sum((row["金额"] for row in report_rows), Decimal("0"))
    message = (
        f"generated: {output_path}; assets={len(report_rows)}; "
        f"total={format_decimal(total_cost)}; latest={latest_path}"
    )
    write_log(args.log, message)
    print(message)
    return 0


def parse_run_date(value: str | None) -> date:
    if not value:
        return date.today()
    return datetime.strptime(value, "%Y-%m-%d").date()


def is_last_day(day: date) -> bool:
    return (day + timedelta(days=1)).month != day.month


def read_env(path: Path) -> dict[str, str]:
    if not path.exists():
        raise FileNotFoundError(f"env file not found: {path}")
    env: dict[str, str] = {}
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def required(env: dict[str, str], key: str) -> str:
    value = env.get(key, "").strip()
    if not value:
        raise RuntimeError(f"missing required env: {key}")
    return value


def fetch_assets(
    *,
    session: requests.Session,
    base_url: str,
    verify_ssl: bool,
    page_size: int,
    max_pages: int,
) -> list[dict[str, Any]]:
    assets: list[dict[str, Any]] = []
    offset = 0

    for _ in range(max_pages):
        data = api_get(
            session,
            base_url,
            "/api/v1/hardware",
            verify_ssl,
            params={
                "limit": page_size,
                "offset": offset,
                "sort": "asset_tag",
                "order": "asc",
            },
        )
        rows = data.get("rows") or []
        if not rows:
            break

        assets.extend(rows)
        total = int(data.get("total") or len(assets))
        offset += page_size
        if offset >= total:
            break

    return assets


def api_get(
    session: requests.Session,
    base_url: str,
    path: str,
    verify_ssl: bool,
    params: dict[str, Any] | None = None,
) -> dict[str, Any]:
    url = urljoin(base_url, path.lstrip("/"))
    response = session.get(url, params=params, timeout=45, verify=verify_ssl)
    if response.status_code == 401:
        raise RuntimeError("Snipe-IT API returned 401. Check SNIPEIT_API_TOKEN.")
    if response.status_code == 403:
        raise RuntimeError("Snipe-IT API returned 403. Token user lacks permission.")
    if response.status_code >= 400:
        raise RuntimeError(f"Snipe-IT API returned HTTP {response.status_code}: {response.text[:300]}")
    return response.json()


def normalize_asset(
    row: dict[str, Any],
    base_url: str,
    rental_prefix: str,
    owned_prefix: str,
) -> dict[str, Any]:
    asset_id = row.get("id")
    asset_tag = text(row.get("asset_tag"))
    model = row.get("model") or {}
    assigned_to = row.get("assigned_to") or {}

    return {
        "序号": 0,
        "资产编号": asset_tag,
        "资产名称": text(row.get("name")),
        "资产来源": classify_ownership(asset_tag, rental_prefix, owned_prefix),
        "分类": nested_name(row.get("category")),
        "型号": nested_name(model),
        "型号编码": text(model.get("model_number") if isinstance(model, dict) else ""),
        "序列号": text(row.get("serial")),
        "当前状态": nested_name(row.get("status_label")),
        "使用人/归属": nested_name(assigned_to),
        "部门": nested_name(assigned_to.get("department") if isinstance(assigned_to, dict) else None),
        "公司": nested_name(row.get("company")),
        "位置": nested_name(row.get("location")) or nested_name(row.get("rtd_location")),
        "供应商": nested_name(row.get("supplier")),
        "采购日期": date_text(row.get("purchase_date")),
        "订单号": text(row.get("order_number")),
        "金额": money(row.get("purchase_cost")),
        "资产链接": urljoin(base_url, f"hardware/{asset_id}") if asset_id else "",
    }


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return nested_name(value)
    return str(value).strip()


def nested_name(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, dict):
        for key in ("name", "full_name", "username", "asset_tag", "serial"):
            item = value.get(key)
            if item:
                return str(item).strip()
        return ""
    return str(value).strip()


def date_text(value: Any) -> str:
    if isinstance(value, dict):
        return text(value.get("date") or value.get("formatted") or value.get("datetime"))
    return text(value)


def money(value: Any) -> Decimal:
    if isinstance(value, dict):
        value = value.get("amount") or value.get("value") or value.get("formatted")
    raw = text(value)
    if not raw:
        return Decimal("0")
    cleaned = re.sub(r"[^0-9.\-]", "", raw.replace(",", ""))
    if cleaned in {"", "-", ".", "-."}:
        return Decimal("0")
    try:
        return Decimal(cleaned).quantize(Decimal("0.01"))
    except InvalidOperation:
        return Decimal("0")


def classify_ownership(asset_tag: str, rental_prefix: str, owned_prefix: str) -> str:
    normalized = asset_tag.upper()
    if rental_prefix and normalized.startswith(rental_prefix):
        return "租赁"
    if owned_prefix and normalized.startswith(owned_prefix):
        return "自购"
    return "未识别"


def natural_sort_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value or "")]


def create_workbook(
    rows: list[dict[str, Any]],
    output_path: Path,
    base_url: str,
    run_date: date,
    rental_prefix: str,
    owned_prefix: str,
) -> None:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "费用清单"
    build_detail_sheet(detail, rows)

    summary = workbook.create_sheet("费用汇总")
    build_summary_sheet(summary, rows, base_url, run_date)

    notes = workbook.create_sheet("生成说明")
    build_notes_sheet(notes, rows, base_url, run_date, rental_prefix, owned_prefix)

    workbook.save(output_path)


def build_detail_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    sheet.append(DETAIL_HEADERS)
    for index, row in enumerate(rows, start=1):
        row["序号"] = index
        sheet.append([row[header] for header in DETAIL_HEADERS])

    apply_sheet_style(sheet, len(rows) + 1, len(DETAIL_HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    amount_col = DETAIL_HEADERS.index("金额") + 1
    link_col = DETAIL_HEADERS.index("资产链接") + 1
    for row_number in range(2, len(rows) + 2):
        sheet.cell(row_number, amount_col).number_format = '#,##0.00'
        link_cell = sheet.cell(row_number, link_col)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"

    add_table(sheet, "MonthlyFeeDetail")
    set_widths(
        sheet,
        {
            "A": 8,
            "B": 24,
            "C": 18,
            "D": 10,
            "E": 16,
            "F": 24,
            "G": 18,
            "H": 24,
            "I": 12,
            "J": 20,
            "K": 16,
            "L": 16,
            "M": 16,
            "N": 18,
            "O": 14,
            "P": 22,
            "Q": 12,
            "R": 42,
        },
    )


def build_summary_sheet(sheet: Any, rows: list[dict[str, Any]], base_url: str, run_date: date) -> None:
    total_cost = sum((row["金额"] for row in rows), Decimal("0"))
    zero_cost = sum(1 for row in rows if row["金额"] == 0)

    sheet["A1"] = f"{run_date.strftime('%Y-%m')} 资产费用汇总"
    sheet["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=PRIMARY)
    sheet.merge_cells("A1:D1")

    summary_rows = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("资产总数", len(rows)),
        ("费用合计", float(total_cost)),
        ("金额为空/0 的资产", zero_cost),
        ("数据来源", base_url.rstrip("/")),
    ]
    for index, item in enumerate(summary_rows, start=3):
        sheet.cell(index, 1).value = item[0]
        sheet.cell(index, 2).value = item[1]
        sheet.cell(index, 1).font = Font(name=FONT_NAME, bold=True, color=PRIMARY_DARK)
        if item[0] == "费用合计":
            sheet.cell(index, 2).number_format = '#,##0.00'

    row_cursor = 10
    row_cursor = add_group_summary(sheet, row_cursor, "按资产来源", rows, "资产来源")
    row_cursor = add_group_summary(sheet, row_cursor + 2, "按分类", rows, "分类")
    row_cursor = add_group_summary(sheet, row_cursor + 2, "按当前状态", rows, "当前状态")
    add_group_summary(sheet, row_cursor + 2, "按公司", rows, "公司")

    set_widths(sheet, {"A": 26, "B": 16, "C": 16, "D": 18})
    sheet.freeze_panes = "A3"


def add_group_summary(sheet: Any, start_row: int, title: str, rows: list[dict[str, Any]], key: str) -> int:
    grouped: dict[str, dict[str, Decimal | int]] = defaultdict(lambda: {"count": 0, "amount": Decimal("0")})
    for row in rows:
        label = row.get(key) or "未填写"
        grouped[str(label)]["count"] += 1
        grouped[str(label)]["amount"] += row["金额"]

    sheet.cell(start_row, 1).value = title
    sheet.cell(start_row, 1).font = Font(name=FONT_NAME, size=12, bold=True, color=PRIMARY_DARK)
    sheet.cell(start_row + 1, 1).value = key
    sheet.cell(start_row + 1, 2).value = "资产数"
    sheet.cell(start_row + 1, 3).value = "金额合计"
    sheet.cell(start_row + 1, 4).value = "占比"
    for col in range(1, 5):
        header = sheet.cell(start_row + 1, col)
        header.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor=PRIMARY)
        header.alignment = Alignment(horizontal="center")

    total_amount = sum((row["金额"] for row in rows), Decimal("0"))
    current = start_row + 2
    for label, item in sorted(grouped.items(), key=lambda entry: (-entry[1]["amount"], entry[0])):
        amount = item["amount"]
        ratio = float(amount / total_amount) if total_amount else 0
        sheet.cell(current, 1).value = label
        sheet.cell(current, 2).value = item["count"]
        sheet.cell(current, 3).value = float(amount)
        sheet.cell(current, 3).number_format = '#,##0.00'
        sheet.cell(current, 4).value = ratio
        sheet.cell(current, 4).number_format = "0.00%"
        current += 1

    apply_sheet_style(sheet, current - 1, 4, start_row=start_row + 1)
    return current


def build_notes_sheet(
    sheet: Any,
    rows: list[dict[str, Any]],
    base_url: str,
    run_date: date,
    rental_prefix: str,
    owned_prefix: str,
) -> None:
    lines = [
        ("报表月份", run_date.strftime("%Y-%m")),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("数据来源", base_url.rstrip("/")),
        ("费用口径", "Snipe-IT 资产字段 purchase_cost，对应页面中的采购价格/金额。"),
        ("资产范围", "通过 /api/v1/hardware 拉取当前可见资产，不包含已被 API 隐藏的删除数据。"),
        (
            "来源识别",
            f"资产编号以 {rental_prefix or '[未配置]'} 开头识别为租赁，"
            f"以 {owned_prefix or '[未配置]'} 开头识别为自购，其余为未识别。",
        ),
        ("自动执行", "计划任务每天运行一次，只有每月最后一天才真正生成文件。"),
    ]
    for row_number, (key, value) in enumerate(lines, start=1):
        sheet.cell(row_number, 1).value = key
        sheet.cell(row_number, 2).value = value
    apply_sheet_style(sheet, len(lines), 2)
    set_widths(sheet, {"A": 18, "B": 96})

    if any(row["金额"] == 0 for row in rows):
        row_number = len(lines) + 2
        sheet.cell(row_number, 1).value = "提醒"
        sheet.cell(row_number, 2).value = "存在金额为空或 0 的资产，请确认 Snipe-IT 中是否已维护采购价格/费用。"
        sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor=WARNING)
        sheet.cell(row_number, 2).fill = PatternFill("solid", fgColor=WARNING)


def apply_sheet_style(sheet: Any, max_row: int, max_col: int, start_row: int = 1) -> None:
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(min_row=start_row, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=10, color="1F2933")
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for cell in sheet[start_row]:
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number in range(start_row + 1, max_row + 1):
        fill = PatternFill("solid", fgColor=PALE_BLUE if row_number % 2 == 0 else "FFFFFF")
        for col in range(1, max_col + 1):
            sheet.cell(row_number, col).fill = fill


def add_table(sheet: Any, name: str) -> None:
    if sheet.max_row < 2:
        return
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def set_widths(sheet: Any, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows():
        sheet.row_dimensions[row[0].row].height = 22


def format_decimal(value: Decimal) -> str:
    return f"{value:,.2f}"


def write_log(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {message}"
    with path.open("a", encoding="utf-8") as handle:
        handle.write(line + os.linesep)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        write_log(DEFAULT_LOG_PATH, f"failed: {exc}")
        print(f"failed: {exc}", file=sys.stderr)
        raise
