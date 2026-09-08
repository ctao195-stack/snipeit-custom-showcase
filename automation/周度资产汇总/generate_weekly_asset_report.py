#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate a weekly Snipe-IT asset summary workbook."""

from __future__ import annotations

import argparse
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE_PATH = Path(r"D:\snipe-it\asset_activity.xlsx")
REPORT_DIR = Path(r"D:\资产周度汇总")
OUTPUT_PATH = REPORT_DIR / "资产周度汇总.xlsx"
ARCHIVE_DIR = REPORT_DIR
STATUS_LABELS = [
    "入库",
    "领用",
    "借用中",
    "归还",
    "退租",
    "维修中",
    "维修完成已重新入库",
    "已换新",
]
DELETE_LABEL = "删除资产"
FALLBACK_LABEL = "未分类"
PRIMARY = "1F4E78"
PRIMARY_DARK = "17365D"
ACCENT = "5B9BD5"
SUCCESS = "70AD47"
WARNING = "F4B183"
DANGER = "C00000"
LIGHT_BLUE = "EAF3F8"
PALE_BLUE = "F6FAFD"
LIGHT_GRAY = "F3F6F8"
MID_GRAY = "D9E2F3"
TEXT_DARK = "1F2933"
BORDER_GRAY = "D9E2EC"
FONT_NAME = "Microsoft YaHei"

COL_SYNC_TIME = 0
COL_LOG_ID = 1
COL_ACTION_TIME = 2
COL_ACTION = 3
COL_OPERATION_STATUS = 4
COL_CURRENT_STATUS = 5
COL_CATEGORY = 6
COL_INTERNAL_ID = 7
COL_ASSET_TAG = 8
COL_ASSET_NAME = 9
COL_SERIAL = 10
COL_OPERATOR = 11
COL_TARGET = 12
COL_NOTE = 13
COL_META = 14

DETAIL_HEADERS = [
    "同步时间",
    "日志ID",
    "操作时间",
    "操作动作",
    "操作后状态",
    "同步时当前状态",
    "资产分类",
    "Snipe-IT内部ID",
    "资产编号",
    "资产名称",
    "序列号",
    "操作人",
    "对象/领用人",
    "备注",
    "变更详情(JSON)",
]

BRIEF_HEADERS = [
    "操作时间",
    "日志ID",
    "操作动作",
    "操作后状态",
    "资产分类",
    "资产编号",
    "资产名称",
    "对象/领用人",
    "备注",
]


@dataclass
class ActivityRow:
    values: list[Any]
    action_time: datetime
    action: str
    bucket: str
    category: str
    asset_tag: str
    target: str
    note: str

    @property
    def log_id(self) -> str:
        return text_at(self.values, COL_LOG_ID)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate weekly asset report.")
    parser.add_argument("--source", type=Path, default=SOURCE_PATH)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--archive-dir", type=Path, default=ARCHIVE_DIR)
    parser.add_argument(
        "--end",
        help="Report end time, for example 2026-05-22 16:00:00. Defaults to now.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    end_at = parse_end(args.end) if args.end else datetime.now()
    start_at = week_start(end_at)
    previous_start_at = start_at - timedelta(days=7)
    previous_end_at = end_at - timedelta(days=7)
    week_label = iso_week_label(end_at)

    rows = load_activity_rows(args.source, start_at, end_at)
    previous_rows = load_activity_rows(args.source, previous_start_at, previous_end_at)
    all_rows = load_activity_rows(args.source, None, end_at)
    workbook = build_report(
        rows,
        previous_rows,
        all_rows,
        start_at,
        end_at,
        previous_start_at,
        previous_end_at,
        week_label,
    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)

    args.archive_dir.mkdir(parents=True, exist_ok=True)
    archive_path = args.archive_dir / f"{period_file_label(start_at, end_at, week_label)}-资产周度汇总.xlsx"
    shutil.copy2(args.output, archive_path)

    print(
        f"Generated {args.output}; archived {archive_path}; "
        f"period={start_at:%Y-%m-%d %H:%M:%S}..{end_at:%Y-%m-%d %H:%M:%S}; rows={len(rows)}"
    )
    return 0


def parse_end(value: str) -> datetime:
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    raise SystemExit("Invalid --end format. Use YYYY-MM-DD HH:MM[:SS].")


def week_start(end_at: datetime) -> datetime:
    monday = end_at.date() - timedelta(days=end_at.weekday())
    return datetime.combine(monday, time.min)


def iso_week_label(value: datetime) -> str:
    year, week, _ = value.isocalendar()
    return f"{year}-W{week:02d}"


def period_file_label(start_at: datetime, end_at: datetime, week_label: str) -> str:
    return f"{week_label}_{start_at:%Y%m%d}-{end_at:%Y%m%d}"


def load_activity_rows(
    source: Path, start_at: datetime | None, end_at: datetime
) -> list[ActivityRow]:
    if not source.exists():
        raise SystemExit(f"Source workbook not found: {source}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    header_map = {
        safe_text(sheet.cell(1, col).value): col - 1
        for col in range(1, sheet.max_column + 1)
        if safe_text(sheet.cell(1, col).value)
    }
    rows: list[ActivityRow] = []

    max_col = max(sheet.max_column, len(DETAIL_HEADERS))
    for raw in sheet.iter_rows(min_row=2, max_col=max_col, values_only=True):
        values = canonical_activity_values(list(raw), header_map)
        action_time = parse_datetime(values[COL_ACTION_TIME] if len(values) > COL_ACTION_TIME else None)
        if not action_time or action_time > end_at:
            continue
        if start_at and action_time < start_at:
            continue

        action = text_at(values, COL_ACTION)
        status = text_at(values, COL_OPERATION_STATUS)
        asset_tag = text_at(values, COL_ASSET_TAG)
        if not asset_tag:
            continue
        if not looks_like_asset_row(values, action, status):
            continue
        bucket = report_bucket(action, status)
        rows.append(
            ActivityRow(
                values=values,
                action_time=action_time,
                action=action,
                bucket=bucket,
                category=text_at(values, COL_CATEGORY) or "未填写",
                asset_tag=asset_tag,
                target=text_at(values, COL_TARGET) or "未填写",
                note=text_at(values, COL_NOTE),
            )
        )

    workbook.close()
    rows.sort(key=lambda row: (row.action_time, row.log_id))
    return rows


def canonical_activity_values(row: list[Any], header_map: dict[str, int]) -> list[Any]:
    def get(*names: str) -> Any:
        for name in names:
            index = header_map.get(name)
            if index is not None and index < len(row):
                return row[index]
        return ""

    if not header_map:
        values = list(row[: len(DETAIL_HEADERS)])
        while len(values) < len(DETAIL_HEADERS):
            values.append("")
        return values

    return [
        get("同步时间"),
        get("日志ID", "_日志ID"),
        get("操作时间"),
        get("操作动作", "操作类型", "操作类型(中文)"),
        get("操作后状态", "状态标签", "资产状态", "当前资产状态"),
        get("同步时当前状态", "_同步时当前状态", "当前资产状态", "资产状态"),
        get("资产分类", "资产类型"),
        get("Snipe-IT内部ID", "_Snipe-IT内部ID", "资产ID"),
        get("资产编号"),
        get("资产名称"),
        get("序列号"),
        get("操作人"),
        get("对象/领用人", "使用人/对象"),
        get("备注"),
        get("变更详情(JSON)", "_变更详情(JSON)"),
    ]


def parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = safe_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def report_bucket(action: str, status: str) -> str:
    if is_delete_action(action):
        return DELETE_LABEL
    clean = status.strip()
    return clean if clean in STATUS_LABELS else FALLBACK_LABEL


def looks_like_asset_row(values: list[Any], action: str, status: str) -> bool:
    return bool(
        text_at(values, COL_ASSET_TAG)
        or text_at(values, COL_CATEGORY)
        or status.strip() in STATUS_LABELS
        or is_delete_action(action)
    )


def is_delete_action(value: str) -> bool:
    text = value.strip().lower()
    return text in {"delete", "deleted", "删除", "删除资产"}


def build_report(
    rows: list[ActivityRow],
    previous_rows: list[ActivityRow],
    all_rows: list[ActivityRow],
    start_at: datetime,
    end_at: datetime,
    previous_start_at: datetime,
    previous_end_at: datetime,
    week_label: str,
) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    counts = Counter(row.bucket for row in rows)
    previous_counts = Counter(row.bucket for row in previous_rows)
    add_overview_sheet(
        workbook,
        rows,
        previous_rows,
        counts,
        start_at,
        end_at,
        previous_start_at,
        previous_end_at,
        week_label,
    )
    add_status_sheet(workbook, counts)
    add_weekly_comparison_sheet(workbook, counts, previous_counts)
    add_new_assets_sheet(workbook, rows)
    add_checkout_ranking_sheet(workbook, rows)
    add_disposal_delete_sheet(workbook, rows)
    add_open_repair_sheet(workbook, all_rows, end_at)
    add_category_sheet(workbook, rows)
    add_people_sheet(workbook, rows)
    add_text_sheet(workbook, rows, previous_rows, counts, start_at, end_at, week_label)
    add_detail_sheet(workbook, rows)
    return workbook


def add_overview_sheet(
    workbook: Workbook,
    rows: list[ActivityRow],
    previous_rows: list[ActivityRow],
    counts: Counter[str],
    start_at: datetime,
    end_at: datetime,
    previous_start_at: datetime,
    previous_end_at: datetime,
    week_label: str,
) -> None:
    sheet = workbook.create_sheet("本周总览")
    repair_count = counts.get("维修中", 0) + counts.get("维修完成已重新入库", 0)
    period_text = f"{start_at:%Y-%m-%d %H:%M:%S} 至 {end_at:%Y-%m-%d %H:%M:%S}"
    total_diff = len(rows) - len(previous_rows)

    sheet.merge_cells("A1:H1")
    sheet["A1"] = "资产周度汇总"
    sheet["A1"].font = Font(name=FONT_NAME, size=22, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=PRIMARY_DARK)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 38

    sheet.merge_cells("A2:H2")
    sheet["A2"] = (
        f"统计周：{week_label}    统计周期：{period_text}    "
        f"生成时间：{datetime.now():%Y-%m-%d %H:%M:%S}"
    )
    sheet["A2"].font = Font(name=FONT_NAME, size=10, color=TEXT_DARK)
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 24

    cards = [
        ("总操作数", len(rows), PRIMARY),
        ("入库", counts.get("入库", 0), SUCCESS),
        ("领用", counts.get("领用", 0), ACCENT),
        ("归还", counts.get("归还", 0), "8064A2"),
        ("退租", counts.get("退租", 0), WARNING),
        ("维修相关", repair_count, "A9D18E"),
        ("删除资产", counts.get(DELETE_LABEL, 0), DANGER),
    ]
    card_ranges = ["A4:B6", "C4:D6", "E4:F6", "G4:H6", "A8:B10", "C8:D10", "E8:F10"]
    for (label, value, color), cell_range in zip(cards, card_ranges):
        add_kpi_card(sheet, cell_range, label, value, color)

    sheet["A12"] = "本周要点"
    sheet["A12"].font = Font(name=FONT_NAME, size=14, bold=True, color=TEXT_DARK)
    sheet.merge_cells("A13:H15")
    sheet["A13"] = overview_note(rows, counts, repair_count, total_diff)
    sheet["A13"].font = Font(name=FONT_NAME, size=11, color=TEXT_DARK)
    sheet["A13"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    sheet["A13"].alignment = Alignment(vertical="top", wrap_text=True)
    apply_range_border(sheet, 13, 1, 15, 8)

    sheet["A17"] = "状态分布"
    sheet["A17"].font = Font(name=FONT_NAME, size=14, bold=True, color=TEXT_DARK)
    start_row = 19
    distribution_rows = [
        ["状态", "数量"],
        ["入库", counts.get("入库", 0)],
        ["领用", counts.get("领用", 0)],
        ["借用中", counts.get("借用中", 0)],
        ["归还", counts.get("归还", 0)],
        ["退租", counts.get("退租", 0)],
        ["维修相关", repair_count],
        ["删除资产", counts.get(DELETE_LABEL, 0)],
        ["未分类", counts.get(FALLBACK_LABEL, 0)],
    ]
    for offset, row in enumerate(distribution_rows):
        row_index = start_row + offset
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row_index, col_index).value = value
    style_table_area(sheet, start_row, 1, start_row + len(distribution_rows) - 1, 2)
    add_status_chart(sheet, start_row, len(distribution_rows))

    for col in range(1, 9):
        sheet.column_dimensions[get_column_letter(col)].width = 14


def add_status_sheet(workbook: Workbook, counts: Counter[str]) -> None:
    sheet = workbook.create_sheet("状态统计")
    total = sum(counts.values())
    sheet.append(["状态/类型", "数量", "占比"])
    for label in [*STATUS_LABELS, DELETE_LABEL, FALLBACK_LABEL]:
        value = counts.get(label, 0)
        ratio = value / total if total else 0
        sheet.append([label, value, ratio])
    for cell in sheet["C"][1:]:
        cell.number_format = "0.0%"
    style_table_sheet(sheet)


def add_weekly_comparison_sheet(
    workbook: Workbook, counts: Counter[str], previous_counts: Counter[str]
) -> None:
    sheet = workbook.create_sheet("上周对比")
    sheet.append(["项目", "本周", "上周同期", "变化", "趋势"])
    labels = ["总操作数", *STATUS_LABELS, DELETE_LABEL, FALLBACK_LABEL]
    for label in labels:
        current = sum(counts.values()) if label == "总操作数" else counts.get(label, 0)
        previous = (
            sum(previous_counts.values()) if label == "总操作数" else previous_counts.get(label, 0)
        )
        diff = current - previous
        sheet.append([label, current, previous, diff, trend_text(diff)])
    style_table_sheet(sheet)


def add_new_assets_sheet(workbook: Workbook, rows: list[ActivityRow]) -> None:
    sheet = workbook.create_sheet("本周新增资产")
    sheet.append(BRIEF_HEADERS)
    written = 0
    for row in rows:
        if row.bucket != "入库":
            continue
        sheet.append(brief_row(row))
        written += 1
    if not written:
        sheet.append(["本周无新增资产"])
    style_table_sheet(sheet)


def add_checkout_ranking_sheet(workbook: Workbook, rows: list[ActivityRow]) -> None:
    sheet = workbook.create_sheet("领用人员排行")
    sheet.append(["排名", "对象/领用人", "领用次数", "涉及资产", "最近操作时间"])
    grouped: dict[str, list[ActivityRow]] = defaultdict(list)
    for row in rows:
        if row.bucket in {"领用", "借用中"}:
            grouped[row.target].append(row)

    ranking = sorted(grouped.items(), key=lambda item: (-len(item[1]), item[0]))
    for index, (person, person_rows) in enumerate(ranking, start=1):
        assets = sorted({row.asset_tag or text_at(row.values, COL_ASSET_NAME) for row in person_rows})
        latest = max(row.action_time for row in person_rows)
        sheet.append([index, person, len(person_rows), "、".join(assets[:8]), latest])
    if not ranking:
        sheet.append(["本周无领用记录"])
    style_table_sheet(sheet)


def add_disposal_delete_sheet(workbook: Workbook, rows: list[ActivityRow]) -> None:
    sheet = workbook.create_sheet("退租删除明细")
    sheet.append(BRIEF_HEADERS)
    written = 0
    for row in rows:
        if row.bucket not in {"退租", DELETE_LABEL}:
            continue
        sheet.append(brief_row(row))
        written += 1
    if not written:
        sheet.append(["本周无退租或删除资产记录"])
    style_table_sheet(sheet)


def add_open_repair_sheet(workbook: Workbook, all_rows: list[ActivityRow], end_at: datetime) -> None:
    sheet = workbook.create_sheet("维修未完成")
    sheet.append([*BRIEF_HEADERS, "截至时间"])
    latest_by_asset: dict[str, ActivityRow] = {}
    for row in all_rows:
        key = row.asset_tag or text_at(row.values, COL_INTERNAL_ID)
        if not key:
            continue
        old = latest_by_asset.get(key)
        if old is None or row.action_time >= old.action_time:
            latest_by_asset[key] = row

    open_repairs = [
        row for row in latest_by_asset.values() if row.bucket == "维修中"
    ]
    open_repairs.sort(key=lambda row: (row.action_time, row.asset_tag))
    for row in open_repairs:
        sheet.append([*brief_row(row), f"{end_at:%Y-%m-%d %H:%M:%S}"])
    if not open_repairs:
        sheet.append(["当前无未完成维修资产"])
    style_table_sheet(sheet)


def add_category_sheet(workbook: Workbook, rows: list[ActivityRow]) -> None:
    sheet = workbook.create_sheet("资产分类统计")
    headers = ["资产分类", "总数", *STATUS_LABELS, DELETE_LABEL, FALLBACK_LABEL]
    sheet.append(headers)
    grouped: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        grouped[row.category][row.bucket] += 1
    for category, counter in sorted(
        grouped.items(), key=lambda item: (-sum(item[1].values()), item[0])
    ):
        sheet.append(
            [
                category,
                sum(counter.values()),
                *[counter.get(label, 0) for label in [*STATUS_LABELS, DELETE_LABEL, FALLBACK_LABEL]],
            ]
        )
    style_table_sheet(sheet)


def add_people_sheet(workbook: Workbook, rows: list[ActivityRow]) -> None:
    sheet = workbook.create_sheet("人员统计")
    headers = ["对象/领用人", "总数", "领用", "借用中", "归还", "退租", "删除资产"]
    sheet.append(headers)
    grouped: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        grouped[row.target][row.bucket] += 1
    for person, counter in sorted(
        grouped.items(), key=lambda item: (-sum(item[1].values()), item[0])
    ):
        sheet.append(
            [
                person,
                sum(counter.values()),
                counter.get("领用", 0),
                counter.get("借用中", 0),
                counter.get("归还", 0),
                counter.get("退租", 0),
                counter.get(DELETE_LABEL, 0),
            ]
        )
    style_table_sheet(sheet)


def add_text_sheet(
    workbook: Workbook,
    rows: list[ActivityRow],
    previous_rows: list[ActivityRow],
    counts: Counter[str],
    start_at: datetime,
    end_at: datetime,
    week_label: str,
) -> None:
    sheet = workbook.create_sheet("周报文字")
    top_categories = category_summary(rows)
    repair_count = counts.get("维修中", 0) + counts.get("维修完成已重新入库", 0)
    total_diff = len(rows) - len(previous_rows)
    diff_phrase = "持平" if total_diff == 0 else f"{'增加' if total_diff > 0 else '减少'} {abs(total_diff)} 次"
    text = (
        f"{week_label} 资产系统周度统计周期为 {start_at:%Y-%m-%d %H:%M} 至 "
        f"{end_at:%Y-%m-%d %H:%M}。本周共记录资产操作 {len(rows)} 次，其中入库 "
        f"{counts.get('入库', 0)} 次、领用 {counts.get('领用', 0)} 次、借用中 "
        f"{counts.get('借用中', 0)} 次、归还 {counts.get('归还', 0)} 次、退租 "
        f"{counts.get('退租', 0)} 次、维修相关 {repair_count} 次、删除资产 "
        f"{counts.get(DELETE_LABEL, 0)} 次，较上周同期{diff_phrase}。"
        f"资产流转主要集中在{top_categories}。"
    )
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "周报文字"
    sheet.merge_cells("A2:C7")
    sheet["A2"] = text
    sheet["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=PRIMARY_DARK)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"].font = Font(name=FONT_NAME, size=12, color=TEXT_DARK)
    sheet["A2"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    for column in "ABC":
        sheet.column_dimensions[column].width = 36
    for row_index in range(2, 8):
        sheet.row_dimensions[row_index].height = 28
    apply_range_border(sheet, 1, 1, 7, 3)


def add_detail_sheet(workbook: Workbook, rows: list[ActivityRow]) -> None:
    sheet = workbook.create_sheet("本周明细")
    sheet.append(DETAIL_HEADERS)
    for row in rows:
        sheet.append(pad_values(row.values))
    style_table_sheet(sheet)


def category_summary(rows: list[ActivityRow]) -> str:
    counter = Counter(row.category for row in rows if row.category and row.category != "未填写")
    if not counter:
        return "暂无明显集中分类"
    names = [name for name, _ in counter.most_common(3)]
    return "、".join(names)


def overview_note(
    rows: list[ActivityRow],
    counts: Counter[str],
    repair_count: int,
    total_diff: int,
) -> str:
    diff_phrase = "与上周同期持平" if total_diff == 0 else (
        f"较上周同期{'增加' if total_diff > 0 else '减少'} {abs(total_diff)} 次"
    )
    return (
        f"本周资产系统共记录操作 {len(rows)} 次，{diff_phrase}。"
        f"其中领用 {counts.get('领用', 0)} 次、归还 {counts.get('归还', 0)} 次、"
        f"入库 {counts.get('入库', 0)} 次、维修相关 {repair_count} 次。"
        f"资产流转主要集中在{category_summary(rows)}。"
    )


def brief_row(row: ActivityRow) -> list[Any]:
    return [
        row.action_time,
        row.log_id,
        row.action,
        row.bucket,
        row.category,
        row.asset_tag,
        text_at(row.values, COL_ASSET_NAME),
        row.target,
        row.note,
    ]


def trend_text(diff: int) -> str:
    if diff > 0:
        return "增加"
    if diff < 0:
        return "减少"
    return "持平"


def pad_values(values: list[Any]) -> list[Any]:
    padded = list(values[: len(DETAIL_HEADERS)])
    while len(padded) < len(DETAIL_HEADERS):
        padded.append("")
    return padded


def text_at(values: list[Any], index: int) -> str:
    return safe_text(values[index] if len(values) > index else "")


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def style_simple_sheet(sheet: Any) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 55
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT_NAME, color=TEXT_DARK)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet["A"]:
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY_DARK)


def style_table_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY_DARK)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 26
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        fill = PatternFill("solid", fgColor=PALE_BLUE if row_index % 2 == 0 else "FFFFFF")
        for cell in row:
            cell.font = Font(name=FONT_NAME, color=TEXT_DARK)
            cell.fill = fill
            cell.border = border
            if isinstance(cell.value, (int, float)) and cell.column > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in sheet.columns:
        letter = column_cells[0].column_letter
        max_len = max(len(safe_text(cell.value)) for cell in column_cells[:200])
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)


def style_table_area(sheet: Any, min_row: int, min_col: int, max_row: int, max_col: int) -> None:
    if max_row < min_row or max_col < min_col:
        return
    sheet.sheet_view.showGridLines = False
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row_index in range(min_row, max_row + 1):
        is_header = row_index == min_row
        fill_color = PRIMARY_DARK if is_header else (PALE_BLUE if row_index % 2 == 0 else "FFFFFF")
        for col_index in range(min_col, max_col + 1):
            cell = sheet.cell(row_index, col_index)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(
                name=FONT_NAME,
                bold=is_header,
                color="FFFFFF" if is_header else TEXT_DARK,
            )
            cell.alignment = Alignment(
                horizontal="center" if is_header or isinstance(cell.value, (int, float)) else "left",
                vertical="center" if is_header else "top",
                wrap_text=True,
            )
    sheet.row_dimensions[min_row].height = 26


def add_status_chart(sheet: Any, start_row: int, row_count: int) -> None:
    if row_count <= 1:
        return
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "状态分布"
    chart.y_axis.title = "数量"
    chart.x_axis.title = "状态"
    chart.height = 8
    chart.width = 16
    data = Reference(sheet, min_col=2, min_row=start_row, max_row=start_row + row_count - 1)
    categories = Reference(sheet, min_col=1, min_row=start_row + 1, max_row=start_row + row_count - 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend = None
    sheet.add_chart(chart, "D19")


def add_kpi_card(sheet: Any, cell_range: str, label: str, value: int, color: str) -> None:
    start, end = cell_range.split(":")
    start_col = column_index(start)
    start_row = row_index(start)
    end_col = column_index(end)
    end_row = row_index(end)
    sheet.merge_cells(cell_range)
    cell = sheet[start]
    cell.value = f"{label}\n{value}"
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(start_row, end_row + 1):
        sheet.row_dimensions[row].height = 24
    apply_range_border(sheet, start_row, start_col, end_row, end_col)


def style_overview_detail(sheet: Any, header_row: int) -> None:
    sheet.sheet_view.showGridLines = False
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(min_row=header_row, max_row=sheet.max_row, max_col=3):
        for cell in row:
            cell.font = Font(name=FONT_NAME, color=TEXT_DARK)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[header_row]:
        if cell.column > 3:
            continue
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index in range(header_row + 1, sheet.max_row + 1):
        fill = PatternFill("solid", fgColor=PALE_BLUE if row_index % 2 == 0 else "FFFFFF")
        for col_index in range(1, 4):
            sheet.cell(row_index, col_index).fill = fill
    sheet.freeze_panes = f"A{header_row + 1}"


def apply_range_border(sheet: Any, start_row: int, start_col: int, end_row: int, end_col: int) -> None:
    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(
        min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col
    ):
        for cell in row:
            cell.border = border


def column_index(cell_ref: str) -> int:
    letters = "".join(char for char in cell_ref if char.isalpha()).upper()
    value = 0
    for char in letters:
        value = value * 26 + (ord(char) - 64)
    return value


def row_index(cell_ref: str) -> int:
    digits = "".join(char for char in cell_ref if char.isdigit())
    return int(digits)


if __name__ == "__main__":
    raise SystemExit(main())
