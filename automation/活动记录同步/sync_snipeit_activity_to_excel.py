#!/usr/bin/env python3
"""Sync Snipe-IT activity logs into an Excel workbook.

This script is intentionally external to Snipe-IT. It talks to the remote
Snipe-IT API, pulls recent activity rows, de-duplicates them, and appends new
rows to an .xlsx file.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urljoin


HEADERS = [
    "操作时间",
    "状态标签",
    "操作动作",
    "资产分类",
    "资产编号",
    "资产名称",
    "序列号",
    "使用人/对象",
    "操作人",
    "备注",
    "同步时间",
    "_日志ID",
    "_Snipe-IT内部ID",
    "_变更详情(JSON)",
    "_同步时当前状态",
]
VISIBLE_COLUMN_COUNT = 11

COL_ACTION_TIME = 0
COL_STATUS = 1
COL_ACTION = 2
COL_CATEGORY = 3
COL_ASSET_TAG = 4
COL_ASSET_NAME = 5
COL_SERIAL = 6
COL_TARGET = 7
COL_OPERATOR = 8
COL_NOTE = 9
COL_SYNC_TIME = 10
COL_LOG_ID = 11
COL_INTERNAL_ID = 12
COL_META = 13
COL_CURRENT_STATUS = 14

OLD_FULL_HEADERS = [
    "同步时间",
    "日志ID",
    "操作时间",
    "操作动作",
    "操作后状态",
    "同步时当前状态",
    "资产分类",
    "Snipe-IT内部ID",
    "资产编号",
    "资产名称",
    "序列号",
    "操作人",
    "对象/领用人",
    "备注",
    "变更详情(JSON)",
]

OLD_CURRENT_STATUS_HEADERS = [
    "同步时间",
    "日志ID",
    "操作时间",
    "操作动作",
    "当前资产状态",
    "资产分类",
    "Snipe-IT内部ID",
    "资产编号",
    "资产名称",
    "序列号",
    "操作人",
    "对象/领用人",
    "备注",
    "变更详情(JSON)",
]

STATUS_LABELS = [
    "借用中",
    "入库",
    "已换新",
    "归还",
    "维修中",
    "维修完成已重新入库",
    "退租",
    "领用",
]
FALLBACK_STATUS_LABEL = "未分类"
DELETE_WORKBOOK_LABEL = "删除资产"
STATUS_COLORS = {
    "入库": "D9EAD3",
    "领用": "DDEBF7",
    "借用中": "E4DFEC",
    "归还": "DDEFE8",
    "退租": "FCE4D6",
    "维修中": "FFF2CC",
    "维修完成已重新入库": "D9EAD3",
    "已换新": "E2F0D9",
    DELETE_WORKBOOK_LABEL: "E7E6E6",
    FALLBACK_STATUS_LABEL: "F3F6F8",
}

FONT_NAME = "Microsoft YaHei"
PRIMARY = "1F4E78"
PRIMARY_DARK = "17365D"
LIGHT_BLUE = "EAF3F8"
PALE_BLUE = "F6FAFD"
BORDER_GRAY = "D9E2EC"
TEXT_DARK = "1F2933"

OLD_COMBINED_ASSET_HEADERS = [
    "同步时间",
    "日志ID",
    "操作时间",
    "操作类型",
    "操作类型(中文)",
    "资产类型",
    "资产ID",
    "资产标签/名称",
    "操作人",
    "对象/领用人",
    "备注",
    "变更详情(JSON)",
]

OLD_SPLIT_ASSET_HEADERS = [
    "同步时间",
    "日志ID",
    "操作时间",
    "操作类型",
    "操作类型(中文)",
    "资产分类",
    "Snipe-IT内部ID",
    "资产编号",
    "资产名称",
    "序列号",
    "操作人",
    "对象/领用人",
    "备注",
    "变更详情(JSON)",
]

OLD_ACTION_STATUS_HEADERS = [
    "同步时间",
    "日志ID",
    "操作时间",
    "操作动作",
    "资产状态",
    "资产分类",
    "Snipe-IT内部ID",
    "资产编号",
    "资产名称",
    "序列号",
    "操作人",
    "对象/领用人",
    "备注",
    "变更详情(JSON)",
]

ACTION_LABELS = {
    "create": "入库/创建",
    "checkout": "借出/领用",
    "checkin from": "归还",
    "checkin": "归还",
    "update": "更新",
    "delete": "删除",
    "audit": "盘点",
    "accepted": "确认接收",
    "declined": "拒绝接收",
    "requested": "申请",
    "uploaded": "上传文件",
    "新增": "入库/新增",
    "借出": "借出/领用",
    "领用": "借出/领用",
    "归还": "归还",
    "归还自": "归还",
    "更新": "更新",
    "删除": "删除",
    "盘点": "盘点",
}


class DependencyError(RuntimeError):
    """Raised when an optional runtime dependency is missing."""


class SnipeItApiError(RuntimeError):
    """Raised for Snipe-IT API responses that need watcher-level handling."""

    def __init__(
        self,
        message: str,
        status_code: int | None = None,
        retry_after: int | None = None,
    ) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.retry_after = retry_after


@dataclass(frozen=True)
class Config:
    base_url: str
    api_token: str
    excel_path: Path
    status_dir: Path
    state_path: Path
    item_type: str
    action_types: set[str]
    page_size: int
    max_pages: int
    verify_ssl: bool
    sheet_name: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Sync remote Snipe-IT activity logs into an Excel workbook."
    )
    parser.add_argument(
        "--env",
        default=".env.snipeit-excel",
        help="Path to env file. Defaults to .env.snipeit-excel",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Fetch and show how many rows would be appended without writing Excel.",
    )
    parser.add_argument(
        "--test-api",
        action="store_true",
        help="Only verify the Snipe-IT API token and connection.",
    )
    parser.add_argument(
        "--rebuild-status-files",
        action="store_true",
        help="Rebuild per-status Excel files from the main workbook.",
    )
    return parser.parse_args()


def load_env_file(path: Path) -> None:
    if not path.exists():
        return

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip().lstrip("\ufeff")
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)


def get_config() -> Config:
    base_url = required_env("SNIPEIT_URL").rstrip("/") + "/"
    api_token = required_env("SNIPEIT_API_TOKEN")
    excel_path = Path(os.getenv("SNIPEIT_EXCEL_PATH", "./snipeit_asset_activity.xlsx"))
    status_dir_value = os.getenv("SNIPEIT_STATUS_DIR", "").strip()
    status_dir = Path(status_dir_value) if status_dir_value else excel_path.parent
    action_types = {
        item.strip().lower()
        for item in os.getenv("SNIPEIT_ACTION_TYPES", "").split(",")
        if item.strip()
    }

    return Config(
        base_url=base_url,
        api_token=api_token,
        excel_path=excel_path,
        status_dir=status_dir,
        state_path=Path(os.getenv("SNIPEIT_STATE_PATH", "./snipeit_excel_sync_state.json")),
        item_type=os.getenv("SNIPEIT_ITEM_TYPE", "asset").strip(),
        action_types=action_types,
        page_size=int(os.getenv("SNIPEIT_PAGE_SIZE", "200")),
        max_pages=int(os.getenv("SNIPEIT_MAX_PAGES", "20")),
        verify_ssl=parse_bool(os.getenv("SNIPEIT_VERIFY_SSL", "true")),
        sheet_name=os.getenv("SNIPEIT_SHEET_NAME", "操作记录"),
    )


def required_env(key: str) -> str:
    value = os.getenv(key, "").strip()
    if not value:
        raise SystemExit(f"Missing required environment variable: {key}")
    return value


def parse_bool(value: str) -> bool:
    return value.strip().lower() not in {"0", "false", "no", "off"}


def require_requests() -> Any:
    try:
        import requests
    except ModuleNotFoundError as exc:
        raise DependencyError(
            "Missing dependency: requests. Run: pip install -r requirements-snipeit-excel.txt"
        ) from exc
    return requests


def require_openpyxl_workbook() -> tuple[Any, Any]:
    try:
        from openpyxl import Workbook, load_workbook
    except ModuleNotFoundError as exc:
        raise DependencyError(
            "Missing dependency: openpyxl. Run: pip install -r requirements-snipeit-excel.txt"
        ) from exc
    return Workbook, load_workbook


def require_openpyxl_styles() -> tuple[Any, Any, Any]:
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
    except ModuleNotFoundError as exc:
        raise DependencyError(
            "Missing dependency: openpyxl. Run: pip install -r requirements-snipeit-excel.txt"
        ) from exc
    return Alignment, Font, PatternFill


def require_openpyxl_tables() -> tuple[Any, Any]:
    try:
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ModuleNotFoundError as exc:
        raise DependencyError(
            "Missing dependency: openpyxl. Run: pip install -r requirements-snipeit-excel.txt"
        ) from exc
    return Table, TableStyleInfo


def api_session(config: Config, requests_module: Any) -> Any:
    session = requests_module.Session()
    session.headers.update(
        {
            "Accept": "application/json",
            "Authorization": f"Bearer {config.api_token}",
        }
    )
    return session


def api_get(
    session: Any, config: Config, path: str, params: dict[str, Any] | None = None
) -> Any:
    url = urljoin(config.base_url, path.lstrip("/"))
    response = session.get(url, params=params, verify=config.verify_ssl, timeout=30)
    if response.status_code == 401:
        raise SnipeItApiError(
            "Snipe-IT API returned 401. Check SNIPEIT_API_TOKEN.",
            status_code=401,
        )
    if response.status_code == 403:
        raise SnipeItApiError(
            "Snipe-IT API returned 403. The token user may lack permission.",
            status_code=403,
        )
    if response.status_code == 429:
        raise SnipeItApiError(
            "Snipe-IT API returned 429 Too Many Requests.",
            status_code=429,
            retry_after=parse_retry_after(response.headers.get("Retry-After")),
        )
    response.raise_for_status()
    return response.json()


def parse_retry_after(value: Any) -> int | None:
    try:
        seconds = int(safe_text(value).strip())
    except ValueError:
        return None
    return seconds if seconds > 0 else None


def test_api(session: Any, config: Config) -> None:
    data = api_get(session, config, "/api/v1/users/me")
    name = data.get("name") or data.get("username") or data.get("email") or "unknown user"
    print(f"API OK: authenticated as {name}")


def load_state(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def save_state(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    state = {
        "last_seen_key": row_key(row),
        "last_seen_log_id": row_id(row),
        "last_seen_created_at": created_at_text(row),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def row_id(row: dict[str, Any]) -> str:
    value = row.get("id") or row.get("log_id") or row.get("action_id")
    return "" if value is None else str(value)


def row_key(row: dict[str, Any]) -> str:
    created = created_at_text(row)
    log_id = row_id(row).zfill(12)
    if created or log_id.strip("0"):
        return f"{created}|{log_id}"

    fallback = "|".join(
        [
            safe_text(row.get("action_type")),
            display_object(row.get("item")),
            safe_text(row.get("note")),
        ]
    )
    return fallback


def created_at_text(row: dict[str, Any]) -> str:
    value = row.get("created_at")
    if isinstance(value, dict):
        value = value.get("datetime") or value.get("formatted") or value.get("date")
    return safe_text(value)


def fetch_activity_rows(
    session: Any, config: Config, last_seen_key: str
) -> list[dict[str, Any]]:
    collected: list[dict[str, Any]] = []
    stop = False

    for page in range(config.max_pages):
        params: dict[str, Any] = {
            "limit": config.page_size,
            "offset": page * config.page_size,
            "sort": "created_at",
            "order": "desc",
        }
        if config.item_type:
            params["item_type"] = config.item_type

        data = api_get(session, config, "/api/v1/reports/activity", params)
        rows = data.get("rows", data if isinstance(data, list) else [])
        if not rows:
            break

        for row in rows:
            if not isinstance(row, dict):
                continue
            current_key = row_key(row)
            if last_seen_key and current_key <= last_seen_key:
                stop = True
                continue

            action_type = safe_text(row.get("action_type")).lower()
            if config.action_types and action_type not in config.action_types:
                continue

            collected.append(row)

        if stop or len(rows) < config.page_size:
            break

    collected.sort(key=row_key)
    return collected


def workbook_and_sheet(config: Config) -> tuple[Any, Any]:
    Workbook, load_workbook = require_openpyxl_workbook()

    if config.excel_path.exists():
        workbook = load_workbook(config.excel_path)
        sheet = workbook[config.sheet_name] if config.sheet_name in workbook.sheetnames else workbook.create_sheet(config.sheet_name)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = config.sheet_name

    schema_changed = migrate_existing_sheet(sheet)
    ensure_header(sheet)
    setattr(sheet, "_snipeit_schema_changed", schema_changed)
    return workbook, sheet


def migrate_existing_sheet(sheet: Any) -> bool:
    current = [
        sheet.cell(1, col).value for col in range(1, len(HEADERS) + 1)
    ]
    if current == HEADERS:
        return False

    old_full = [
        sheet.cell(1, col).value for col in range(1, len(OLD_FULL_HEADERS) + 1)
    ]
    old_combined = [
        sheet.cell(1, col).value for col in range(1, len(OLD_COMBINED_ASSET_HEADERS) + 1)
    ]
    old_split = [
        sheet.cell(1, col).value for col in range(1, len(OLD_SPLIT_ASSET_HEADERS) + 1)
    ]
    old_current = [
        sheet.cell(1, col).value for col in range(1, len(OLD_CURRENT_STATUS_HEADERS) + 1)
    ]
    if old_full == OLD_FULL_HEADERS:
        max_cols = len(OLD_FULL_HEADERS)
        migrator = migrate_old_full_excel_row
    elif old_combined == OLD_COMBINED_ASSET_HEADERS:
        max_cols = len(OLD_COMBINED_ASSET_HEADERS)
        migrator = migrate_old_combined_excel_row
    elif old_split == OLD_SPLIT_ASSET_HEADERS:
        max_cols = len(OLD_SPLIT_ASSET_HEADERS)
        migrator = migrate_old_split_excel_row
    elif old_split == OLD_ACTION_STATUS_HEADERS:
        max_cols = len(OLD_ACTION_STATUS_HEADERS)
        migrator = migrate_old_action_status_excel_row
    elif old_current == OLD_CURRENT_STATUS_HEADERS:
        max_cols = len(OLD_CURRENT_STATUS_HEADERS)
        migrator = migrate_old_current_status_excel_row
    else:
        return False

    migrated_rows = []
    for row in sheet.iter_rows(
        min_row=2, max_row=sheet.max_row, max_col=max_cols, values_only=True
    ):
        if not any(value not in (None, "") for value in row):
            continue
        migrated_rows.append(migrator(list(row)))

    clear_tables(sheet)
    sheet.delete_rows(1, sheet.max_row)
    sheet.append(HEADERS)
    for row in migrated_rows:
        sheet.append(row)
    return True


def migrate_old_full_excel_row(row: list[Any]) -> list[Any]:
    while len(row) < len(OLD_FULL_HEADERS):
        row.append("")
    return clean_display_row_from_canonical(row)


def clean_display_row_from_canonical(row: list[Any]) -> list[Any]:
    while len(row) < len(OLD_FULL_HEADERS):
        row.append("")
    return [
        row[2],
        row[4],
        row[3],
        row[6],
        row[8],
        row[9],
        row[10],
        row[12],
        row[11],
        row[13],
        row[0],
        row[1],
        row[7],
        row[14],
        row[5],
    ]


def migrate_old_combined_excel_row(row: list[Any]) -> list[Any]:
    while len(row) < len(OLD_COMBINED_ASSET_HEADERS):
        row.append("")

    asset = parse_asset_display(safe_text(row[7]))
    category = safe_text(row[5]) or asset["category"]
    action = normalize_action_label(row[4] or row[3])
    meta = row[11]
    operation_status = infer_operation_status(action, meta, "", asset["asset_tag"])

    return clean_display_row_from_canonical([
        row[0],
        row[1],
        row[2],
        action,
        operation_status,
        "",
        category,
        row[6],
        asset["asset_tag"],
        asset["name"],
        asset["serial"],
        row[8],
        row[9],
        row[10],
        meta,
    ])


def migrate_old_split_excel_row(row: list[Any]) -> list[Any]:
    while len(row) < len(OLD_SPLIT_ASSET_HEADERS):
        row.append("")

    action = normalize_action_label(row[4] or row[3])
    meta = row[13]
    return clean_display_row_from_canonical([
        row[0],
        row[1],
        row[2],
        action,
        infer_operation_status(action, meta, "", row[7]),
        "",
        row[5],
        row[6],
        row[7],
        row[8],
        row[9],
        row[10],
        row[11],
        row[12],
        meta,
    ])


def migrate_old_action_status_excel_row(row: list[Any]) -> list[Any]:
    while len(row) < len(OLD_ACTION_STATUS_HEADERS):
        row.append("")

    status = clean_status_text(row[4])
    return clean_display_row_from_canonical([
        row[0],
        row[1],
        row[2],
        row[3],
        status_from_meta(row[13]) or status,
        status,
        row[5],
        row[6],
        row[7],
        row[8],
        row[9],
        row[10],
        row[11],
        row[12],
        row[13],
    ])


def migrate_old_current_status_excel_row(row: list[Any]) -> list[Any]:
    while len(row) < len(OLD_CURRENT_STATUS_HEADERS):
        row.append("")

    action = row[3]
    current_status = clean_status_text(row[4])
    meta = row[13]
    return clean_display_row_from_canonical([
        row[0],
        row[1],
        row[2],
        action,
        infer_operation_status(action, meta, current_status, row[7]),
        current_status,
        row[5],
        row[6],
        row[7],
        row[8],
        row[9],
        row[10],
        row[11],
        row[12],
        meta,
    ])


def clear_tables(sheet: Any) -> None:
    for name in list(sheet.tables.keys()):
        del sheet.tables[name]


def ensure_header(sheet: Any) -> None:
    Alignment, Font, PatternFill = require_openpyxl_styles()

    if sheet.max_row == 1 and all(sheet.cell(1, col).value is None for col in range(1, len(HEADERS) + 1)):
        for col, header in enumerate(HEADERS, start=1):
            sheet.cell(1, col).value = header

    existing = [sheet.cell(1, col).value for col in range(1, len(HEADERS) + 1)]
    if existing != HEADERS:
        for col, header in enumerate(HEADERS, start=1):
            if sheet.cell(1, col).value is None:
                sheet.cell(1, col).value = header

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(HEADERS) + 1):
        cell = sheet.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"


def existing_log_ids(sheet: Any) -> set[str]:
    ids: set[str] = set()
    log_id_col = COL_LOG_ID + 1
    for row in sheet.iter_rows(min_row=2, min_col=log_id_col, max_col=log_id_col, values_only=True):
        if row[0] not in (None, ""):
            ids.add(str(row[0]))
    return ids


def excel_rows_by_log_ids(sheet: Any, log_ids: set[str]) -> list[list[Any]]:
    if not log_ids:
        return []

    rows: list[list[Any]] = []
    for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        log_id = safe_text(row[COL_LOG_ID] if len(row) > COL_LOG_ID else "")
        if log_id in log_ids:
            rows.append(list(row))
    return rows


def all_excel_rows(sheet: Any) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        if any(value not in (None, "") for value in row):
            rows.append(list(row))
    return rows


def append_status_workbooks(config: Config, excel_rows: list[list[Any]]) -> int:
    if not excel_rows:
        return 0

    grouped: dict[str, list[list[Any]]] = {}
    for row in excel_rows:
        if is_delete_excel_row(row):
            continue
        status = row_status(row)
        grouped.setdefault(status_bucket(status), []).append(row)

    written = 0
    for status, rows in grouped.items():
        workbook, sheet = status_workbook_and_sheet(config, status)
        known_ids = existing_log_ids(sheet)
        schema_changed = getattr(sheet, "_snipeit_schema_changed", False)
        changed = False
        for row in rows:
            log_id = safe_text(row[COL_LOG_ID] if len(row) > COL_LOG_ID else "")
            if log_id and log_id in known_ids:
                continue
            sheet.append(pad_excel_row(row))
            if log_id:
                known_ids.add(log_id)
            written += 1
            changed = True

        if changed or schema_changed:
            config.status_dir.mkdir(parents=True, exist_ok=True)
            style_sheet(sheet)
            workbook.save(status_workbook_path(config, status))

    return written


def append_delete_workbook(config: Config, excel_rows: list[list[Any]]) -> int:
    delete_rows = [row for row in excel_rows if is_delete_excel_row(row)]
    if not delete_rows:
        return 0

    workbook, sheet = routed_workbook_and_sheet(config, delete_workbook_path(config))
    known_ids = existing_log_ids(sheet)
    schema_changed = getattr(sheet, "_snipeit_schema_changed", False)
    changed = False
    written = 0

    for row in delete_rows:
        log_id = safe_text(row[COL_LOG_ID] if len(row) > COL_LOG_ID else "")
        if log_id and log_id in known_ids:
            continue
        sheet.append(pad_excel_row(row))
        if log_id:
            known_ids.add(log_id)
        changed = True
        written += 1

    if changed or schema_changed:
        config.status_dir.mkdir(parents=True, exist_ok=True)
        style_sheet(sheet)
        workbook.save(delete_workbook_path(config))

    return written


def rebuild_status_workbooks(config: Config, sheet: Any) -> int:
    rows = all_excel_rows(sheet)
    config.status_dir.mkdir(parents=True, exist_ok=True)

    for status in [*STATUS_LABELS, FALLBACK_STATUS_LABEL]:
        path = status_workbook_path(config, status)
        if path.exists():
            path.unlink()
        workbook, status_sheet = status_workbook_and_sheet(config, status)
        style_sheet(status_sheet)
        workbook.save(path)

    return append_status_workbooks(config, rows)


def rebuild_delete_workbook(config: Config, sheet: Any) -> int:
    path = delete_workbook_path(config)
    config.status_dir.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()
    workbook, delete_sheet = routed_workbook_and_sheet(config, path)
    style_sheet(delete_sheet)
    workbook.save(path)
    return append_delete_workbook(config, all_excel_rows(sheet))


def status_workbook_and_sheet(config: Config, status: str) -> tuple[Any, Any]:
    return routed_workbook_and_sheet(config, status_workbook_path(config, status))


def routed_workbook_and_sheet(config: Config, path: Path) -> tuple[Any, Any]:
    Workbook, load_workbook = require_openpyxl_workbook()
    if path.exists():
        workbook = load_workbook(path)
        sheet = (
            workbook[config.sheet_name]
            if config.sheet_name in workbook.sheetnames
            else workbook.create_sheet(config.sheet_name)
        )
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = config.sheet_name

    schema_changed = migrate_existing_sheet(sheet)
    ensure_header(sheet)
    setattr(sheet, "_snipeit_schema_changed", schema_changed)
    return workbook, sheet


def status_workbook_path(config: Config, status: str) -> Path:
    return config.status_dir / f"{safe_filename(status_bucket(status))}.xlsx"


def delete_workbook_path(config: Config) -> Path:
    return config.status_dir / f"{safe_filename(DELETE_WORKBOOK_LABEL)}.xlsx"


def row_status(row: list[Any]) -> str:
    return clean_status_text(row[COL_STATUS] if len(row) > COL_STATUS else "")


def row_action(row: list[Any]) -> str:
    return safe_text(row[COL_ACTION] if len(row) > COL_ACTION else "")


def is_delete_excel_row(row: list[Any]) -> bool:
    return is_delete_action(row_action(row))


def is_delete_action(value: Any) -> bool:
    text = safe_text(value).strip().lower()
    return text in {"delete", "deleted", "删除", "删除资产"}


def status_bucket(status: Any) -> str:
    clean = clean_status_text(status)
    return clean if clean in STATUS_LABELS else FALLBACK_STATUS_LABEL


def safe_filename(value: str) -> str:
    text = re.sub(r'[<>:"/\\|?*]+', "_", value).strip()
    return text or FALLBACK_STATUS_LABEL


def pad_excel_row(row: list[Any]) -> list[Any]:
    values = list(row[: len(HEADERS)])
    while len(values) < len(HEADERS):
        values.append("")
    return values


def append_rows(
    sheet: Any,
    rows: list[dict[str, Any]],
    known_ids: set[str],
    session: Any | None = None,
    config: Config | None = None,
) -> list[dict[str, Any]]:
    appended: list[dict[str, Any]] = []
    sync_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    status_cache: dict[str, str] = {}

    for row in rows:
        log_id = row_id(row)
        if log_id and log_id in known_ids:
            continue
        if config and not activity_matches_item_type(row, config.item_type):
            continue

        sheet.append(to_excel_row(row, sync_time, session, config, status_cache))
        if log_id:
            known_ids.add(log_id)
        appended.append(row)

    return appended


def to_excel_row(
    row: dict[str, Any],
    sync_time: str,
    session: Any | None = None,
    config: Config | None = None,
    status_cache: dict[str, str] | None = None,
) -> list[str]:
    action_type = safe_text(row.get("action_type"))
    action_label = normalize_action_label(action_type)
    item = row.get("item")
    asset = asset_details(item, row)
    target = row.get("target") or row.get("checkedout_to") or row.get("assigned_to")
    meta = row.get("log_meta") or row.get("metadata") or row.get("changes") or ""
    current_status = status_label_text(item) or current_asset_status(
        item, session, config, status_cache
    )
    operation_status = infer_operation_status(
        action_type,
        meta,
        current_status,
        asset["asset_tag"],
    )

    return [
        created_at_text(row),
        operation_status,
        action_label,
        asset["category"],
        asset["asset_tag"],
        asset["name"],
        asset["serial"],
        display_object(target),
        display_object(row.get("admin") or row.get("user")),
        safe_text(row.get("note") or row.get("notes")),
        sync_time,
        row_id(row),
        item_id(item),
        json_text(meta),
        current_status,
    ]


def normalize_action_label(value: Any) -> str:
    action_type = safe_text(value)
    return ACTION_LABELS.get(action_type.lower(), action_type)


def infer_operation_status(
    action_type: Any,
    meta: Any,
    current_status: Any = "",
    asset_tag: Any = "",
) -> str:
    if is_checkin_action(action_type):
        return "归还"

    metadata_status = status_from_meta(meta)
    if metadata_status:
        return metadata_status

    if is_create_action(action_type) and safe_text(asset_tag).strip():
        return "入库"

    return clean_status_text(current_status)


def is_create_action(value: Any) -> bool:
    action = safe_text(value).strip().lower()
    return action in {"create", "新增", "入库/新增", "入库/创建"}


def is_checkin_action(value: Any) -> bool:
    action = safe_text(value).strip().lower()
    return action in {"checkin from", "checkin", "归还", "归还自"}


def current_asset_status(
    item: Any,
    session: Any | None,
    config: Config | None,
    cache: dict[str, str] | None = None,
) -> str:
    if session is None or config is None:
        return ""

    asset_id = item_id(item)
    if not asset_id:
        return ""
    if cache is not None and asset_id in cache:
        return cache[asset_id]

    try:
        data = api_get(session, config, f"/api/v1/hardware/{asset_id}")
    except SnipeItApiError as exc:
        if exc.status_code in {401, 403, 429}:
            raise
        data = {}
    except Exception:
        data = {}

    status = clean_status_text(status_label_text(data))
    if cache is not None:
        cache[asset_id] = status
    return status


def status_label_text(data: Any) -> str:
    if not isinstance(data, dict):
        return ""

    value = (
        data.get("status_label")
        or data.get("statusLabel")
        or data.get("status")
        or data.get("status_label_name")
    )
    if isinstance(value, dict):
        return clean_status_text(value.get("name") or value.get("label") or value.get("title"))
    return clean_status_text(value)


def refresh_asset_statuses(sheet: Any, session: Any, config: Config) -> bool:
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    try:
        status_col = (
            headers.index("_同步时当前状态") + 1
            if "_同步时当前状态" in headers
            else headers.index("同步时当前状态") + 1
        )
        asset_id_col = (
            headers.index("_Snipe-IT内部ID") + 1
            if "_Snipe-IT内部ID" in headers
            else headers.index("Snipe-IT内部ID") + 1
        )
    except ValueError:
        return False

    cache: dict[str, str] = {}
    changed = False
    for row_index in range(2, sheet.max_row + 1):
        asset_id = safe_text(sheet.cell(row_index, asset_id_col).value)
        if not asset_id:
            continue
        if asset_id not in cache:
            try:
                cache[asset_id] = status_label_text(
                    api_get(session, config, f"/api/v1/hardware/{asset_id}")
                )
            except Exception:
                cache[asset_id] = ""
        status = cache[asset_id]
        if status and sheet.cell(row_index, status_col).value != status:
            sheet.cell(row_index, status_col).value = status
            changed = True

    return changed


def status_from_meta(meta: Any) -> str:
    if not meta:
        return ""
    if isinstance(meta, str):
        try:
            meta = json.loads(meta)
        except json.JSONDecodeError:
            return parse_status_value(meta)
    if not isinstance(meta, dict):
        return ""

    for key, value in meta.items():
        key_text = safe_text(key).lower()
        if "状态" not in key_text and "status" not in key_text:
            continue
        if isinstance(value, dict):
            return parse_status_value(value.get("new") or value.get("after") or value.get("to"))
        return parse_status_value(value)
    return ""


def parse_status_value(value: Any) -> str:
    text = safe_text(value).strip()
    if not text:
        return ""
    match = re.search(r"\]\s*(?P<status>.+)$", text)
    if match:
        return clean_status_text(match.group("status"))
    return clean_status_text(text)


def clean_status_text(value: Any) -> str:
    text = safe_text(value).strip()
    text = re.sub(r"[.。…]+$", "", text).strip()
    return text


def asset_details(item: Any, row: dict[str, Any]) -> dict[str, str]:
    display = display_object(item)
    parsed_display = parse_asset_display(display)

    if not isinstance(item, dict):
        return parsed_display

    raw_name = safe_text(item.get("name"))
    parsed_name = parse_asset_display(raw_name)
    asset_tag = safe_text(item.get("asset_tag") or item.get("assetTag"))
    serial = safe_text(item.get("serial"))

    return {
        "category": parsed_name["category"] or parsed_display["category"] or item_type_text(row),
        "asset_tag": asset_tag or parsed_name["asset_tag"] or parsed_display["asset_tag"],
        "name": parsed_name["name"] or parsed_display["name"] or raw_name,
        "serial": serial or parsed_name["serial"] or parsed_display["serial"],
    }


def parse_asset_display(value: str) -> dict[str, str]:
    text = safe_text(value).strip()
    result = {"category": "", "asset_tag": "", "name": text, "serial": ""}
    if not text:
        return result

    serial_match = re.search(r"\s*/\s*SN[:：]\s*(?P<serial>.+?)\s*$", text, re.IGNORECASE)
    if serial_match:
        result["serial"] = serial_match.group("serial").strip()
        text = text[: serial_match.start()].strip()

    asset_match = re.match(
        r"^\s*(?P<category>[^()]+?)\s*\((?P<asset_tag>[^)]+)\)\s*-\s*(?P<name>.+?)\s*$",
        text,
    )
    if asset_match:
        result["category"] = asset_match.group("category").strip()
        result["asset_tag"] = asset_match.group("asset_tag").strip()
        result["name"] = asset_match.group("name").strip()
        return result

    tag_match = re.match(
        r"^\s*(?P<category>[^()]+?)\s*\((?P<asset_tag>[^)]+)\)\s*$",
        text,
    )
    if tag_match:
        result["category"] = tag_match.group("category").strip()
        result["asset_tag"] = tag_match.group("asset_tag").strip()
        result["name"] = ""
        return result

    result["name"] = text
    return result


def item_type_text(row: dict[str, Any]) -> str:
    value = row.get("item_type") or row.get("itemType")
    if isinstance(value, dict):
        return safe_text(value.get("name") or value.get("type"))
    return safe_text(value)


def activity_matches_item_type(row: dict[str, Any], desired_type: str) -> bool:
    desired = normalize_item_type(desired_type)
    if not desired:
        return True

    item = row.get("item")
    candidates = [item_type_text(row)]
    if isinstance(item, dict):
        candidates.extend(
            [
                safe_text(item.get("type")),
                safe_text(item.get("item_type") or item.get("itemType")),
            ]
        )

    normalized_candidates = {
        normalize_item_type(candidate) for candidate in candidates if safe_text(candidate)
    }
    if normalized_candidates:
        return desired in normalized_candidates

    if desired == "asset":
        return bool(asset_details(item, row).get("asset_tag"))

    return False


def normalize_item_type(value: Any) -> str:
    text = safe_text(value).strip().lower()
    aliases = {
        "assets": "asset",
        "hardware": "asset",
        "users": "user",
        "consumables": "consumable",
    }
    return aliases.get(text, text)


def item_id(value: Any) -> str:
    if isinstance(value, dict):
        raw_id = value.get("id")
        return "" if raw_id is None else str(raw_id)
    return ""


def display_object(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, dict):
        parts = []
        asset_tag = value.get("asset_tag") or value.get("assetTag")
        name = value.get("name") or value.get("username") or value.get("email")
        serial = value.get("serial")
        if asset_tag:
            parts.append(str(asset_tag))
        if name and str(name) not in parts:
            parts.append(str(name))
        if serial:
            parts.append(f"SN:{serial}")
        if parts:
            return " / ".join(parts)
        raw_id = value.get("id")
        return "" if raw_id is None else f"ID:{raw_id}"
    if isinstance(value, list):
        return ", ".join(display_object(item) for item in value)
    return json_text(value)


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


def json_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def style_sheet(sheet: Any) -> None:
    Alignment, Font, PatternFill = require_openpyxl_styles()
    try:
        from openpyxl.styles import Border, Side
    except ModuleNotFoundError as exc:
        raise DependencyError(
            "Missing dependency: openpyxl. Run: pip install -r requirements-snipeit-excel.txt"
        ) from exc

    sheet.sheet_view.showGridLines = False
    widths = {
        "A": 20,
        "B": 20,
        "C": 12,
        "D": 14,
        "E": 24,
        "F": 28,
        "G": 20,
        "H": 16,
        "I": 14,
        "J": 34,
        "K": 20,
        "L": 14,
        "M": 16,
        "N": 45,
        "O": 18,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    for index in range(VISIBLE_COLUMN_COUNT + 1, len(HEADERS) + 1):
        sheet.column_dimensions[column_letter(index)].hidden = True

    thin = Side(style="thin", color=BORDER_GRAY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor=PRIMARY_DARK)
    header_font = Font(name=FONT_NAME, color="FFFFFF", bold=True)

    for col in range(1, len(HEADERS) + 1):
        cell = sheet.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28

    for row_index, row in enumerate(
        sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(HEADERS)),
        start=2,
    ):
        row_fill = PatternFill(fill_type="solid", fgColor=PALE_BLUE if row_index % 2 else "FFFFFF")
        for cell in row:
            cell.font = Font(name=FONT_NAME, color=TEXT_DARK)
            cell.border = border
            cell.fill = row_fill
            if cell.column in {COL_STATUS + 1, COL_ACTION + 1}:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        status = safe_text(sheet.cell(row_index, COL_STATUS + 1).value)
        status_color = STATUS_COLORS.get(status)
        if status_color:
            status_cell = sheet.cell(row_index, COL_STATUS + 1)
            status_cell.fill = PatternFill(fill_type="solid", fgColor=status_color)
            status_cell.font = Font(name=FONT_NAME, color=TEXT_DARK, bold=True)
        sheet.row_dimensions[row_index].height = 28

    ensure_table(sheet)


def ensure_table(sheet: Any) -> None:
    Table, TableStyleInfo = require_openpyxl_tables()

    if sheet.max_row < 1:
        return
    clear_tables(sheet)
    ref = f"A1:{column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"
    table_name = "SnipeITActivity"
    table = Table(displayName=table_name, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    sheet.add_table(table)


def column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def main() -> int:
    args = parse_args()
    load_env_file(Path(args.env))
    config = get_config()
    requests_module = require_requests()
    session = api_session(config, requests_module)

    try:
        if args.test_api:
            test_api(session, config)
            return 0
    except SnipeItApiError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except requests_module.RequestException as exc:
        print(f"Snipe-IT API request failed: {exc}", file=sys.stderr)
        return 1

    excel_existed = config.excel_path.exists()
    workbook, sheet = workbook_and_sheet(config)
    schema_changed = getattr(sheet, "_snipeit_schema_changed", False)

    if args.rebuild_status_files:
        if schema_changed or not excel_existed:
            config.excel_path.parent.mkdir(parents=True, exist_ok=True)
            style_sheet(sheet)
            workbook.save(config.excel_path)
        written = rebuild_status_workbooks(config, sheet)
        deleted = rebuild_delete_workbook(config, sheet)
        print(
            f"Rebuilt status workbooks in {config.status_dir}; "
            f"routed {written} status row(s), {deleted} delete row(s)."
        )
        return 0

    try:
        state = load_state(config.state_path)
        rows = fetch_activity_rows(session, config, state.get("last_seen_key", ""))
    except SnipeItApiError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except requests_module.RequestException as exc:
        print(f"Snipe-IT API request failed: {exc}", file=sys.stderr)
        return 1

    if args.dry_run:
        print(f"Fetched {len(rows)} new activity row(s). Excel was not modified.")
        return 0

    appended = append_rows(sheet, rows, existing_log_ids(sheet), session, config)
    routed_rows = excel_rows_by_log_ids(sheet, {row_id(row) for row in rows})

    if appended or not excel_existed or schema_changed:
        config.excel_path.parent.mkdir(parents=True, exist_ok=True)
        style_sheet(sheet)
        workbook.save(config.excel_path)

    routed_count = append_status_workbooks(config, routed_rows)
    deleted_count = append_delete_workbook(config, routed_rows)

    if rows:
        save_state(config.state_path, rows[-1])
        print(
            f"Appended {len(appended)} row(s) to {config.excel_path}; "
            f"routed {routed_count} status row(s), "
            f"{deleted_count} delete row(s) to {config.status_dir}"
        )
    elif not excel_existed:
        print(f"Created empty Excel workbook at {config.excel_path}")
    else:
        print("No new Snipe-IT activity rows to append.")

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except DependencyError as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
    except SnipeItApiError as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
    except PermissionError as exc:
        print(
            "Cannot write the Excel file. Close it in Excel or choose another output path: "
            f"{exc}",
            file=sys.stderr,
        )
        raise SystemExit(1)
