#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(r"D:\snipe-it")
SHEET_NAME = "ActivityLog"
FILES = [
    "asset_activity.xlsx",
    "借用中.xlsx",
    "入库.xlsx",
    "已换新.xlsx",
    "归还.xlsx",
    "维修中.xlsx",
    "维修完成已重新入库.xlsx",
    "退租.xlsx",
    "领用.xlsx",
    "未分类.xlsx",
]


def workbook_summary(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
    headers = [sheet.cell(1, col).value for col in range(1, min(sheet.max_column, 15) + 1)]
    return {
        "file": path.name,
        "rows": max(sheet.max_row - 1, 0),
        "columns": sheet.max_column,
        "headers": headers,
    }


def main() -> int:
    summaries = []
    for name in FILES:
        path = BASE / name
        if path.exists():
            summaries.append(workbook_summary(path))
    print(json.dumps(summaries, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
