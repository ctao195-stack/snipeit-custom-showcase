#!/usr/bin/env python3
from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook


def load_sync_module():
    module_path = Path(__file__).resolve().parent / "sync_snipeit_activity_to_excel.py"
    spec = importlib.util.spec_from_file_location("snipeit_excel_sync", module_path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"Cannot load {module_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def main() -> int:
    sync = load_sync_module()

    assert len(sync.HEADERS) == 15
    assert sync.HEADERS[1] == "状态标签"
    assert sync.HEADERS[14] == "_同步时当前状态"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ActivityLog"
    sheet.append(sync.OLD_CURRENT_STATUS_HEADERS)
    sheet.append(
        [
            "2026-05-19 12:00:00",
            "1",
            "2026-05-19 11:59:00",
            "入库/新增",
            "领用",
            "输出设备",
            "1080",
            "RENT-GZ-B-002-01002",
            "AOC",
            "SN1",
            "Test Admin",
            "",
            "",
            "",
        ]
    )
    assert sync.migrate_existing_sheet(sheet)
    assert [sheet.cell(1, col).value for col in range(1, 16)] == sync.HEADERS
    migrated_row = [sheet.cell(2, col).value for col in range(1, 16)]
    assert migrated_row[1] == "入库", migrated_row
    assert migrated_row[14] == "领用", migrated_row

    excel_row = sync.to_excel_row(
        {
            "id": 99,
            "created_at": {"datetime": "2026-05-19 12:05:00"},
            "action_type": "create",
            "item": {
                "id": 1088,
                "asset_tag": "RENT-GZ-B-003-01088",
                "name": "办公设备 (RENT-GZ-B-003-01088) - 笔记本",
                "serial": "SN-TEST",
                "status_label": {"name": "入库"},
            },
            "admin": {"name": "Test Admin"},
        },
        "2026-05-19 12:05:05",
    )
    assert excel_row[1] == "入库", excel_row
    assert excel_row[4] == "RENT-GZ-B-003-01088", excel_row
    assert excel_row[14] == "入库", excel_row

    print("schema logic ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
